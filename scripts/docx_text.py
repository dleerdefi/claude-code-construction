#!/usr/bin/env python3
"""Dump a Word document's text: every paragraph with its style, then every table.

Usage:
    docx_text.py "<file.docx>"
    docx_text.py "<file.docx>" --output "<dir>/template.txt"

Paragraphs print as "[Style Name] text"; tables print row by row with " | "
between cells. Fill-field placeholders such as [DATE] or <<SCOPE>> appear
verbatim, so a template's article structure and its blanks can be read without
writing python-docx code.
"""

import argparse
import sys


def dump(path):
    try:
        import docx
    except ImportError:
        print("ERROR: python-docx is not installed in the construction environment")
        sys.exit(1)
    d = docx.Document(path)
    lines = ["=== paragraphs"]
    for p in d.paragraphs:
        if p.text.strip():
            lines.append(f"[{p.style.name}] {p.text}")
    for i, t in enumerate(d.tables, 1):
        lines.append(f"=== table {i} ({len(t.rows)} rows)")
        for r in t.rows:
            lines.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="Dump a .docx as text with styles and tables")
    parser.add_argument("docx", help="Path to the Word document")
    parser.add_argument("--output", "-o", help="Write to this file instead of stdout")
    args = parser.parse_args()
    try:
        body = dump(args.docx)
    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(body + "\n")
        print(f"OK: {args.output}")
    else:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        print(body)


if __name__ == "__main__":
    main()
