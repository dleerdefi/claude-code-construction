"""qto_to_xlsx.py: pending items stay out of the count and are listed on their own sheet."""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "skills" / "tag-audit-and-takeoff" / "scripts" / "qto_to_xlsx.py"


def qto(instance_statuses, pending_review=None):
    data = {
        "project": {"name": "Test", "generated_at": "2026-10-12T00:00:00"},
        "scope": {"tag_type": "room_tag", "sheets_scanned": ["A-101"]},
        "totals": {"sheet_instances": len(instance_statuses), "derived_instances": 0, "deduplicated": 0,
                   "building_quantity": len(instance_statuses)},
        "completeness": {},
        "type_definitions_applied": [],
        "line_items": [{"element": "Room tag", "designation": "RT", "sheet_instances": len(instance_statuses),
                        "derived_instances": 0, "deduplicated": 0, "building_qty": len(instance_statuses),
                        "instance_details": [{"type": "element_instance", "detection_id": f"d{i}", "sheet": "A-101",
                                              "room": str(100 + i), "status": s} for i, s in enumerate(instance_statuses)],
                        "derived_details": []}],
    }
    if pending_review is not None:
        data["pending_review"] = pending_review
    return data


class PendingSplit(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def export(self, data):
        src = self.dir / "qto.json"
        src.write_text(json.dumps(data), encoding="utf-8")
        out = self.dir / "QTO.xlsx"
        r = subprocess.run([sys.executable, str(SCRIPT), "--data", str(src), "--output", str(out)],
                           capture_output=True, text=True, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
        return r, out

    def test_a_pending_item_inside_the_count_is_refused(self):
        r, out = self.export(qto(["accepted", "pending"]))
        self.assertEqual(r.returncode, 2, r.stdout + r.stderr)
        self.assertIn("pending_review", r.stderr)
        self.assertFalse(out.exists(), "no workbook is written")

    def test_pending_review_is_listed_on_its_own_sheet(self):
        pending = [{"tag_text": "STAIR 2", "tag_type": "room_tag", "sheet": "A-101", "detection_id": "g9",
                    "status": "pending", "note": "pending in AgentCM's review"}]
        r, out = self.export(qto(["accepted", "new"], pending))
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        ws = load_workbook(out)["To Confirm in AgentCM"]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=3)]
        self.assertEqual(rows[0][:2], ["STAIR 2", "room_tag"])
        self.assertEqual(rows[0][4], "g9")

    def test_accepted_approved_and_new_count_and_an_empty_list_says_so(self):
        r, out = self.export(qto(["accepted", "approved", "new"]))
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        wb = load_workbook(out)
        self.assertIn("No pending items.", [c.value for c in wb["To Confirm in AgentCM"]["A"]])


if __name__ == "__main__":
    unittest.main()
