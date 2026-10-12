"""agentcm_schedule.py: AgentCM's table files to schedule_to_xlsx.py's data and the schedules/ingest payload, on
the contract fixture's own two tables (one read from the sheet's text, one ingested)."""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "schedule-extractor" / "scripts"
SCRIPT = SCRIPTS / "agentcm_schedule.py"
TABLES = ROOT / "evals" / "plugin" / "agentcm-mode" / "construction" / "schedules"
sys.path.insert(0, str(SCRIPTS))

import agentcm_schedule as a  # noqa: E402

TEXT, INGESTED = TABLES / "A-601-1.tsv", TABLES / "A-602-1.tsv"
CONTRACT = ("sheetId", "scheduleType", "title", "rowEntityType", "columns", "rows", "sourceMethod", "confidence")


class ReadTable(unittest.TestCase):
    def test_a_table_read_from_the_text_drops_its_title_and_takes_its_header_row(self):
        t = a.read_table(TEXT)
        self.assertEqual((t["sheet"], t["title"], t["schedule_id"]), ("A-601", "DOOR SCHEDULE", None))
        self.assertEqual(t["box"], [0.1, 0.06, 0.44, 0.346])
        self.assertEqual(t["headers"], ["MARK", "SIZE", "DOOR", "FRAME", "HDWR"])
        self.assertEqual(t["rows"][0][:2], ["D101", "3'-0\" x 7'-0\""])
        self.assertTrue(all(r[0].startswith("D") for r in t["rows"]))
        self.assertEqual(t, a.read_table(TEXT, header_row="0.100"))  # the same row, named

    def test_an_ingested_table_keeps_the_drawings_headers_and_its_id(self):
        t = a.read_table(INGESTED)
        self.assertEqual(t["schedule_id"], "ac35b0bc-9b30-4046-bf47-53bf90f698c6")
        self.assertEqual(t["headers"], ["MARK", "WIDTH", "HEIGHT", "TYPE", "GLAZING"])
        self.assertEqual(t["rows"][0], ["A", "2'-0\"", "3'-0\"", "FIXED", "TEMPERED"])

    def test_a_header_row_that_is_not_there_is_refused(self):
        with self.assertRaises(a.TableError):
            a.read_table(TEXT, header_row="0.555")

    def test_two_header_lines_join_per_column_and_a_cell_without_a_header_is_noted(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "M-601-1.tsv"
            p.write_text("# M-601 FAN SCHEDULE: a schedule at 0.1 0.1 0.5 0.3 (x y w h), from the sheet's text. schedule_id=none\n"
                         "y\tx0.1\tx0.2\tx0.3\n0.10\tFAN\tAIR\t\n0.12\tTAG\tCFM\t\n0.14\tEF-1\t400\tsee note\n", encoding="utf-8")
            t = a.read_table(p, header_row="0.10", header_rows=2)
        self.assertEqual(t["headers"], ["FAN TAG", "AIR CFM"])
        self.assertEqual(t["rows"], [["EF-1", "400"]])
        self.assertTrue(any("no header" in n for n in t["notes"]))


class Payload(unittest.TestCase):
    def test_the_payload_has_contract_1s_fields_and_agentcms_values(self):
        body, keyless = a.payload(a.read_table(TEXT), "sheet-id", "door", "door", "hybrid", 0.9)
        self.assertEqual([k for k in CONTRACT if k not in body], [])
        self.assertEqual(keyless, [])
        self.assertEqual(body["rows"][0]["entityIdentifier"], "D101")
        self.assertEqual(set(body["rows"][0]["cells"]), {c["key"] for c in body["columns"]})
        self.assertEqual(body["boundingRegion"], {"x": 0.1, "y": 0.06, "width": 0.44, "height": 0.346})

    def test_an_ingested_table_is_not_ingested_again(self):
        with self.assertRaises(a.TableError) as e:
            a.payload(a.read_table(INGESTED), "sheet-id", "window", "window", "hybrid", 0.9)
        self.assertIn("--schedule-id ac35b0bc", str(e.exception))

    def test_column_keys_are_upper_case_and_unique(self):
        self.assertEqual(a.column_keys(["Hardware Set", "FIRE RATING", "NO.", "No"]), ["HARDWARE_SET", "FIRE_RATING", "NO", "NO_2"])


class SameKey(unittest.TestCase):
    def test_the_workbooks_row_key_and_the_payloads_identifier_agree_when_the_mark_is_not_first(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            tsv, data, out, xlsx = (Path(tmp) / n for n in ("A-601-1.tsv", "data.json", "p.json", "doors.xlsx"))
            tsv.write_text("# A-601 DOOR SCHEDULE: a schedule at 0.1 0.06 0.44 0.3 (x y w h), from the sheet's text. schedule_id=none\n"
                           "y\tx0.1\tx0.2\tx0.3\n0.10\tTYPE\tMARK\tSIZE\n0.12\tA\tD101\t3'-0\"\n0.14\tB\tD102\t3'-6\"\n", encoding="utf-8")
            r = subprocess.run([sys.executable, str(SCRIPT), str(tsv), "--type", "door_schedule", "--data", str(data),
                                "--payload", str(out), "--sheet-id", "s", "--row-entity-type", "door",
                                "--source-method", "hybrid", "--confidence", "0.9"], capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            r = subprocess.run([sys.executable, str(SCRIPTS / "schedule_to_xlsx.py"), "--data", str(data), "--type", "door_schedule",
                                "--sheet", "A-601", "--output", str(xlsx)], capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            ws = load_workbook(xlsx).worksheets[0]
            row_keys = [ws.cell(row=i, column=1).value for i in range(5, 7)]
            body = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(row_keys, ["D101", "D102"])
        self.assertEqual([r["entityIdentifier"] for r in body["rows"]], ["D101", "D102"])
        self.assertEqual(body["scheduleType"], "door")


class Command(unittest.TestCase):
    def run_script(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True,
                              env=dict(os.environ, PYTHONIOENCODING="cp1252"))  # a Windows pipe

    def test_data_and_payload_files_and_agentcms_value_lists(self):
        with tempfile.TemporaryDirectory() as tmp:
            data, out = Path(tmp) / "a601.json", Path(tmp) / "outbox" / "p.json"
            r = self.run_script(str(TEXT), "--type", "door_schedule", "--data", str(data), "--payload", str(out),
                                "--sheet-id", "s", "--row-entity-type", "door", "--source-method", "hybrid", "--confidence", "0.9")
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            rows = json.loads(data.read_text(encoding="utf-8"))
            self.assertEqual(rows[0], ["MARK", "SIZE", "DOOR", "FRAME", "HDWR"])
            self.assertEqual(json.loads(out.read_text(encoding="utf-8"))["scheduleType"], "door")
            r = self.run_script(str(TEXT), "--type", "door", "--data", str(data), "--payload", str(out), "--sheet-id", "s",
                                "--row-entity-type", "door", "--source-method", "agentcm_text", "--confidence", "0.9")
            self.assertEqual(r.returncode, 2)  # not one of AgentCM's source methods
            r = self.run_script(str(INGESTED), "--type", "window", "--data", str(data), "--payload", str(out),
                                "--sheet-id", "s", "--row-entity-type", "window", "--source-method", "hybrid", "--confidence", "0.9")
            self.assertEqual(r.returncode, 2)
            self.assertIn(b"already ingested", r.stdout)


if __name__ == "__main__":
    unittest.main()
