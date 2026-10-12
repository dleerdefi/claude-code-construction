#!/usr/bin/env python3
"""QTO Excel Export — RIGID script.

Claude Code calls this with structured JSON data matching the QTO output
format from references/qto-output-format.md. This script does not make
any judgment calls — it counts and formats the data identically every time.

The count: per line item, confirmed (instances accepted in AgentCM's review), to confirm
(instances waiting there as pending cards or in the outbox, or on a sheet whose file is behind,
plus every derived instance not deduplicated), seen on the drawing, not in AgentCM (status seen:
the drawing shows the tag and AgentCM has no text for it, so it has no card; each with its crop)
and the possible total (all three). The script computes the numbers from the instances itself,
prints one line per item, and refuses a file that lists the same tag twice (same text, one box's
center inside the other, on one sheet), a seen tag without its crop, or the retired pending_review
key. A tag the PE rejected in AgentCM (status rejected) is in no number: it is listed once, apart.
When line items name their group (the schedule whose key column holds the code), it prints the
numbers per group too, "not in any schedule" last with its codes.

Input:  JSON file with project{}, scope{}, line_items[], etc.
Output: Formatted .xlsx workbook with seven sheets:
  - QTO Summary: one row per line item with the four numbers and confidence
  - Instance Detail: every detection record with full provenance
  - Type Definitions: type defs applied and their room targets
  - Completeness: coverage metrics, gap sheets, and issues
  - To Confirm in AgentCM: every instance not yet accepted, with its reason
  - Seen, not in AgentCM: every tag AgentCM has no text for, its box and crop, and how to clear it
  - Rejected in AgentCM: every instance the PE already turned down there, not counted

Usage:
    construction-python qto_to_xlsx.py \\
        --data qto_data.json \\
        --project "Sanibel Fire and Rescue Station 172" \\
        --scope "Room Tags" \\
        --output "QTO_Room_Tags.xlsx"

Exit codes: 0 workbook written; 1 data file missing; 2 refused (no workbook).
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from shared import safe_output_path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter


# ── Color Palette ──────────────────────────────────────────────────

COLORS = {
    "header_bg":     "2C3E50",
    "header_fg":     "FFFFFF",
    "row_even":      "F8F9FA",
    "row_odd":       "FFFFFF",
    "high":          "D4EDDA",
    "medium":        "FFF3CD",
    "low":           "FFE0B2",
    "flagged":       "F8D7DA",
    "high_font":     "155724",
    "medium_font":   "856404",
    "low_font":      "E65100",
    "flagged_font":  "721C24",
    "title_bg":      "1A3C5E",
    "subtitle_bg":   "34495E",
    "accent":        "3498DB",
    "border":        "DEE2E6",
    "dedup_bg":      "E9ECEF",
}

THIN_BORDER = Border(
    left=Side(style="thin", color=COLORS["border"]),
    right=Side(style="thin", color=COLORS["border"]),
    top=Side(style="thin", color=COLORS["border"]),
    bottom=Side(style="thin", color=COLORS["border"]),
)

HEADER_FONT = Font(name="Arial", bold=True, color=COLORS["header_fg"], size=10)
HEADER_FILL = PatternFill("solid", fgColor=COLORS["header_bg"])
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name="Arial", size=10)
WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
CENTER_ALIGN = Alignment(horizontal="center", vertical="top")
RIGHT_ALIGN = Alignment(horizontal="right", vertical="top")

TITLE_FONT = Font(name="Arial", bold=True, color=COLORS["header_fg"], size=14)
TITLE_FILL = PatternFill("solid", fgColor=COLORS["title_bg"])
SUBTITLE_FONT = Font(name="Arial", bold=True, color=COLORS["header_fg"], size=11)
SUBTITLE_FILL = PatternFill("solid", fgColor=COLORS["subtitle_bg"])

CONFIDENCE_STYLES = {
    "high":    (COLORS["high"],    COLORS["high_font"]),
    "medium":  (COLORS["medium"],  COLORS["medium_font"]),
    "low":     (COLORS["low"],     COLORS["low_font"]),
    "flagged": (COLORS["flagged"], COLORS["flagged_font"]),
}


# ── The count ──────────────────────────────────────────────────────

ACCEPTED = {"accepted", "approved"}             # approved: older files
TO_CONFIRM = {"to_confirm", "new", "pending"}   # new, pending: older files
REJECTED = {"rejected"}                         # the PE turned it down in AgentCM: never counted
SEEN = {"seen", "image_only"}                   # on the drawing, no text in AgentCM: no card to accept

CAVEAT_AGENTCM = (
    "Confirmed: tags accepted in AgentCM's review. To confirm: tags waiting in AgentCM's review as pending "
    "cards (made this run or earlier) or in the outbox, tags on a sheet whose file is behind or edited, and "
    "every count derived from a type definition. Seen on the drawing, not in AgentCM: tags the drawing shows "
    "that AgentCM has no text for, so no card; each with its crop. Possible total: all three, what has been "
    "seen so far, not a ceiling."
)
CLEAR_SEEN = (
    "To clear one: in AgentCM, draw the tag's text on the sheet's canvas and tag it. On the next run it shows "
    "as pending or accepted. Each is also in .construction/skills/corrections.tsv (kind tag), once."
)
CAVEAT_FLAT = (
    "No review step in this project (no AgentCM): every tag found is to confirm on the drawings, so none "
    "is confirmed yet. Possible total: what this run saw, not a ceiling."
)


def caveat(data):
    return CAVEAT_FLAT if data.get("scope", {}).get("mode") == "flat_file" else CAVEAT_AGENTCM


def status_of(det):
    """accepted, to_confirm, seen or rejected; None for anything else."""
    s = str(det.get("status", "")).strip().lower()
    if s in ACCEPTED:
        return "accepted"
    if s in TO_CONFIRM:
        return "to_confirm"
    if s in SEEN:
        return "seen"
    if s in REJECTED:
        return "rejected"
    return None


def label(item):
    return " ".join(p for p in (str(item.get("designation", "")).strip(), str(item.get("element", "")).strip()) if p) or "?"


def counts(item):
    """(confirmed, to confirm, seen not in AgentCM, of which derived, possible total) for one line item."""
    statuses = [status_of(d) for d in item.get("instance_details", [])]
    derived = sum(1 for d in item.get("derived_details", []) if not d.get("deduplicated"))
    confirmed = statuses.count("accepted")
    to_confirm = statuses.count("to_confirm") + derived
    seen = statuses.count("seen")
    return confirmed, to_confirm, seen, derived, confirmed + to_confirm + seen


def numbers(c):
    """The line the export prints for a count: the four groups."""
    return f"confirmed {c[0]}, to confirm {c[1]}, seen not in AgentCM {c[2]}, possible {c[4]}"


def all_counts(data):
    per_item = [(item, counts(item)) for item in data.get("line_items", [])]
    total = tuple(sum(c[i] for _, c in per_item) for i in range(5))
    return per_item, total


NOT_IN_SCHEDULE = "not in any schedule"


def groups(data):
    """[(group, (confirmed, to confirm, seen, of which derived, possible), designations)] in the order the groups first
    appear, "not in any schedule" last; empty when no line item names its group. A group is the schedule whose
    key column holds the item's code, as the skill read it from the files: the report's wording, nothing more."""
    per_item, _ = all_counts(data)
    if not any(str(item.get("group", "")).strip() for item, _ in per_item):
        return []
    order, sums, codes = [], {}, {}
    for item, c in per_item:
        g = str(item.get("group", "")).strip() or NOT_IN_SCHEDULE
        g = NOT_IN_SCHEDULE if g.lower() == NOT_IN_SCHEDULE else g
        if g not in sums:
            order.append(g)
            sums[g], codes[g] = (0, 0, 0, 0, 0), []
        sums[g] = tuple(a + b for a, b in zip(sums[g], c))
        codes[g].append(str(item.get("designation", "")).strip() or label(item))
    order.sort(key=lambda g: g == NOT_IN_SCHEDULE)
    return [(g, sums[g], codes[g]) for g in order]


def rejected(data):
    """(line item, instance) for every tag the PE rejected in AgentCM."""
    return with_status(data, "rejected")


def seen(data):
    """(line item, instance) for every tag on the drawing that AgentCM has no text for."""
    return with_status(data, "seen")


def with_status(data, status):
    return [(item, det) for item in data.get("line_items", []) for det in item.get("instance_details", [])
            if status_of(det) == status]


def norm_text(text):
    return " ".join(str(text or "").upper().split())


def box_of(det):
    """The instance's bounding_box as (x, y, width, height), None when absent; ValueError when malformed."""
    b = det.get("bounding_box")
    if b is None:
        return None
    try:
        box = tuple(float(b[k]) for k in ("x", "y", "width", "height"))
    except (TypeError, KeyError, ValueError):
        raise ValueError(f"bounding_box must be {{x, y, width, height}}, got {json.dumps(b)[:80]}")
    return box


def center_inside(a, b):
    cx, cy = a[0] + a[2] / 2, a[1] + a[3] / 2
    return b[0] <= cx <= b[0] + b[2] and b[1] <= cy <= b[1] + b[3]


def problems(data):
    """Everything that stops the export, as messages: an unknown status, a malformed box, a tag listed twice,
    a seen tag without its crop, the retired pending_review key."""
    found = []
    if "pending_review" in data:
        found.append("pending_review is retired: a tag pending in AgentCM's review goes in line_items with "
                     "status to_confirm and source pending_in_agentcm")
    placed = {}  # (sheet, text) -> [(box, where)]
    for item in data.get("line_items", []):
        for det in item.get("instance_details", []):
            where = f"{label(item)}: {det.get('detection_id') or det.get('tag_text') or '?'} on {det.get('sheet', '?')}"
            if status_of(det) is None:
                found.append(f"{where} has status {det.get('status')!r}: use accepted, to_confirm, seen or rejected")
            elif status_of(det) == "seen" and not str(det.get("crop", "")).strip():
                found.append(f"{where} is seen on the drawing, not in AgentCM: it needs its crop")
            try:
                box = box_of(det)
            except ValueError as e:
                found.append(f"{where}: {e}")
                continue
            if box is None:
                continue
            key = (str(det.get("sheet", "")), norm_text(det.get("tag_text") or item.get("designation")))
            for other_box, other in placed.get(key, []):
                if center_inside(box, other_box) or center_inside(other_box, box):
                    found.append(f"the same tag twice: {other} and {where} (same text, one box inside the other)")
            placed.setdefault(key, []).append((box, where))
        for der in item.get("derived_details", []):
            if status_of(der) is None and der.get("status") is not None:
                found.append(f"{label(item)}: derived {der.get('derived_id', '?')} has status {der.get('status')!r}")
    return found


# ── Column Definitions ────────────────────────────────────────────

SUMMARY_COLUMNS = [
    ("Group",            26),
    ("Element",          24),
    ("Designation",      14),
    ("Confirmed",        12),
    ("To Confirm",       12),
    ("Seen, Not in AgentCM", 14),
    ("Possible Total",   14),
    ("Of Which Derived", 14),
    ("Avg Confidence",   14),
    ("Notes",            40),
]

DETAIL_COLUMNS = [
    ("Sheet",          12),
    ("Tag Text",       24),
    ("Designation",    14),
    ("Type",           16),
    ("View Type",      14),
    ("Room",           10),
    ("Source",         20),
    ("Confidence",     12),
    ("Status",         12),
    ("Detection ID",   36),
    ("Element IDs",    40),
]

TYPEDEF_COLUMNS = [
    ("Source Detail",      16),
    ("Parent View Label",  24),
    ("Elements Defined",   40),
    ("Applied To Rooms",   40),
    ("Applied Count",      14),
    ("Derivation Method",  18),
]

COMPLETENESS_COLUMNS = [
    ("Metric",     24),
    ("Value",      40),
]

TO_CONFIRM_COLUMNS = [
    ("Tag Text",     22),
    ("Designation",  14),
    ("Sheet",        10),
    ("Room",         10),
    ("Source",       20),
    ("Ingest",       12),
    ("Detection ID", 38),
    ("Reason",       34),
    ("Note",         50),
]

SEEN_COLUMNS = [
    ("Tag Text",     22),
    ("Designation",  14),
    ("Sheet",        10),
    ("Room",         10),
    ("Box (x y w h)", 24),
    ("Crop",         44),
    ("Corrections",  24),
    ("Note",         50),
]

REJECTED_COLUMNS = [
    ("Tag Text",     22),
    ("Designation",  14),
    ("Sheet",        10),
    ("Room",         10),
    ("Rejected Card", 38),
    ("Note",         50),
]


# ── Helpers ────────────────────────────────────────────────────────

def apply_header(ws, columns, row=1):
    """Write header row and set column widths."""
    for col_idx, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=row, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(columns))}{row}"


def style_body_cell(cell, row_idx, align=None):
    """Apply standard body styling with alternating row shading."""
    cell.font = BODY_FONT
    cell.border = THIN_BORDER
    cell.alignment = align or WRAP_ALIGN
    fill_color = COLORS["row_even"] if row_idx % 2 == 0 else COLORS["row_odd"]
    cell.fill = PatternFill("solid", fgColor=fill_color)


def style_confidence_cell(cell, confidence_val):
    """Color-code a cell based on confidence level."""
    if isinstance(confidence_val, (int, float)):
        if confidence_val >= 0.8:
            level = "high"
        elif confidence_val >= 0.5:
            level = "medium"
        elif confidence_val > 0:
            level = "low"
        else:
            level = "flagged"
    elif isinstance(confidence_val, str):
        level = confidence_val.lower()
    else:
        return

    if level in CONFIDENCE_STYLES:
        bg, fg = CONFIDENCE_STYLES[level]
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(name="Arial", size=10, color=fg, bold=True)


def write_title_row(ws, text, col_count, row=1):
    """Merge and style a title row spanning all columns."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 30
    return row + 1


def write_note_row(ws, text, col_count, row):
    """A merged, wrapped line of explanation under a title."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name="Arial", italic=True, size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 42
    return row + 1


# ── Sheet Builders ─────────────────────────────────────────────────

def build_summary_sheet(ws, data, project_name, scope_name):
    """Build the QTO Summary sheet."""
    ws.title = "QTO Summary"
    per_item, (t_conf, t_todo, t_seen, t_derived, t_possible) = all_counts(data)

    row = write_title_row(ws, f"QTO Summary — {project_name} — {scope_name}", len(SUMMARY_COLUMNS))

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(SUMMARY_COLUMNS))
    cell = ws.cell(row=row, column=1,
                   value=f"Confirmed: {t_conf}  |  To Confirm: {t_todo}  |  Seen, Not in AgentCM: {t_seen}  |  "
                         f"Possible Total: {t_possible}")
    cell.font = SUBTITLE_FONT
    cell.fill = SUBTITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    row += 1
    row = write_note_row(ws, caveat(data), len(SUMMARY_COLUMNS), row)

    apply_header(ws, SUMMARY_COLUMNS, row=row)
    row += 1

    for item, (confirmed, to_confirm, seen_n, derived, possible) in per_item:
        details = item.get("instance_details", [])
        confs = [d.get("confidence", 0) for d in details if isinstance(d.get("confidence"), (int, float))]
        avg_conf = sum(confs) / len(confs) if confs else 0.0
        dedup = sum(1 for d in item.get("derived_details", []) if d.get("deduplicated"))
        notes = f"{dedup} derived dedup'd (direct takes precedence)" if dedup else ""

        group = str(item.get("group", "")).strip() or (NOT_IN_SCHEDULE if groups(data) else "")
        values = [group, item.get("element", ""), item.get("designation", ""), confirmed, to_confirm, seen_n, possible,
                  derived, avg_conf, notes]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            style_body_cell(cell, row)
            if col_idx in (4, 5, 6, 7, 8):
                cell.number_format = "#,##0"
                cell.alignment = RIGHT_ALIGN
            elif col_idx == 9:
                cell.number_format = "0.00"
                cell.alignment = CENTER_ALIGN
                style_confidence_cell(cell, val)
        row += 1

    ws.cell(row=row, column=1, value="TOTAL").font = Font(name="Arial", bold=True, size=10)
    for col_idx, val in [(4, t_conf), (5, t_todo), (6, t_seen), (7, t_possible), (8, t_derived)]:
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.border = THIN_BORDER
        cell.number_format = "#,##0"
        cell.alignment = RIGHT_ALIGN


def build_detail_sheet(ws, data):
    """Build the Instance Detail sheet."""
    ws.title = "Instance Detail"
    apply_header(ws, DETAIL_COLUMNS)
    row = 2

    for item in data.get("line_items", []):
        for det in item.get("instance_details", []):
            values = [
                det.get("sheet", ""),
                det.get("tag_text", item.get("element", "")),
                item.get("designation", ""),
                det.get("type", "element_instance"),
                det.get("view_type", "floor_plan"),
                det.get("room", ""),
                det.get("source", det.get("provenance", "")),
                det.get("confidence", 0),
                status_of(det),
                det.get("detection_id", ""),
                ", ".join(det.get("element_ids", [])) if det.get("element_ids") else "",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                style_body_cell(cell, row)
                if col_idx == 8:
                    cell.number_format = "0.00"
                    cell.alignment = CENTER_ALIGN
                    style_confidence_cell(cell, val)
                elif col_idx == 9:
                    cell.alignment = CENTER_ALIGN
            row += 1

        for der in item.get("derived_details", []):
            values = [
                der.get("target_sheet", ""),
                item.get("element", ""),
                item.get("designation", ""),
                "derived_instance",
                "derived",
                der.get("target_room", ""),
                f"derived from {der.get('source_detail', '')} ({der.get('derivation_method', '')})",
                "",
                "deduplicated" if der.get("deduplicated") else "to_confirm",
                der.get("derived_id", ""),
                "",
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                style_body_cell(cell, row)
                if col_idx == 9:
                    cell.alignment = CENTER_ALIGN

            # Mark deduplicated rows with muted background
            if der.get("deduplicated"):
                for col_idx in range(1, len(DETAIL_COLUMNS) + 1):
                    ws.cell(row=row, column=col_idx).fill = PatternFill("solid", fgColor=COLORS["dedup_bg"])
            row += 1


def build_typedef_sheet(ws, data):
    """Build the Type Definitions sheet."""
    ws.title = "Type Definitions"
    apply_header(ws, TYPEDEF_COLUMNS)
    row = 2

    for td in data.get("type_definitions_applied", []):
        elements_str = "; ".join(
            f"{e.get('tag_text', '')} x{e.get('count_per_instance', 1)}"
            for e in td.get("elements_defined", [])
        )
        rooms_str = ", ".join(str(r) for r in td.get("applied_to_rooms", []))

        values = [
            td.get("source_detail", ""),
            td.get("parent_view_label", ""),
            elements_str,
            rooms_str,
            td.get("applied_count", 0),
            td.get("derivation_method", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            style_body_cell(cell, row)
            if col_idx == 5:
                cell.number_format = "#,##0"
                cell.alignment = CENTER_ALIGN

        row += 1

    if row == 2:
        ws.cell(row=2, column=1, value="No type definitions applied").font = Font(name="Arial", italic=True, size=10)


def build_completeness_sheet(ws, data, project_name, scope_name):
    """Build the Completeness sheet."""
    ws.title = "Completeness"

    row = write_title_row(
        ws,
        f"Completeness Report — {project_name} — {scope_name}",
        len(COMPLETENESS_COLUMNS),
    )
    apply_header(ws, COMPLETENESS_COLUMNS, row=row)
    row += 1

    scope = data.get("scope", {})
    comp = data.get("completeness", {})
    _, (t_conf, t_todo, t_seen, t_derived, t_possible) = all_counts(data)

    metrics = [
        ("Tag Type", scope.get("tag_type", "")),
        ("CSI Division", scope.get("csi_division", "")),
        ("Sheets Scanned", str(len(scope.get("sheets_scanned", [])))),
        ("Sheets with Detections", str(len(scope.get("sheets_with_detections", [])))),
        ("Sheets with Zero", ", ".join(scope.get("sheets_with_zero", [])) or "None"),
        ("", ""),
        ("Confirmed", str(t_conf)),
        ("To Confirm", str(t_todo)),
        ("Of Which Derived", str(t_derived)),
        ("Seen, Not in AgentCM", str(t_seen)),
        ("Possible Total", str(t_possible)),
        ("", ""),
        ("Schedule Reference", comp.get("schedule_reference", "N/A")),
        ("Expected Count", str(comp.get("expected_count", "N/A"))),
        ("Detected Count", str(comp.get("detected_count", "N/A"))),
        ("Coverage", f"{comp.get('coverage_pct', 0):.1f}%" if comp.get("coverage_pct") else "N/A"),
        ("Gap Sheets", ", ".join(comp.get("gap_sheets", [])) or "None"),
        ("Gap Notes", comp.get("gap_notes", "")),
    ]

    for metric, value in metrics:
        if not metric:
            row += 1
            continue
        m_cell = ws.cell(row=row, column=1, value=metric)
        v_cell = ws.cell(row=row, column=2, value=value)
        style_body_cell(m_cell, row)
        style_body_cell(v_cell, row)
        m_cell.font = Font(name="Arial", bold=True, size=10)
        row += 1

    # Issues section
    issues = data.get("issues", [])
    if issues:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = ws.cell(row=row, column=1, value="Issues & Warnings")
        cell.font = SUBTITLE_FONT
        cell.fill = SUBTITLE_FILL
        row += 1

        for issue in issues:
            severity = issue.get("severity", "info").upper()
            msg = issue.get("message", "")
            action = issue.get("suggested_action", "")
            text = f"[{severity}] {msg}"
            if action:
                text += f" — Action: {action}"

            cell = ws.cell(row=row, column=1, value=text)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            style_body_cell(cell, row)

            if severity == "WARNING":
                cell.fill = PatternFill("solid", fgColor=COLORS["medium"])
                cell.font = Font(name="Arial", size=10, color=COLORS["medium_font"])
            elif severity == "ERROR":
                cell.fill = PatternFill("solid", fgColor=COLORS["flagged"])
                cell.font = Font(name="Arial", size=10, color=COLORS["flagged_font"])

            row += 1

    # Generated timestamp
    row += 1
    ts = data.get("project", {}).get("generated_at", datetime.now().isoformat())
    ws.cell(row=row, column=1, value=f"Generated: {ts}").font = Font(
        name="Arial", italic=True, size=9, color="808080",
    )


def build_to_confirm_sheet(ws, data):
    """Every instance not yet accepted, and every derived instance counted: listed for the PE to confirm."""
    ws.title = "To Confirm in AgentCM"
    row = write_title_row(ws, "Not yet accepted: confirm in AgentCM (inside the possible total, not the confirmed count)",
                          len(TO_CONFIRM_COLUMNS))
    row = write_note_row(ws, caveat(data), len(TO_CONFIRM_COLUMNS), row)
    apply_header(ws, TO_CONFIRM_COLUMNS, row=row)
    row += 1
    first = row
    for item in data.get("line_items", []):
        for det in item.get("instance_details", []):
            if status_of(det) != "to_confirm":
                continue
            values = [det.get("tag_text", ""), item.get("designation", ""), det.get("sheet", ""), det.get("room", ""),
                      det.get("source", ""), det.get("ingest", ""), det.get("detection_id", ""), det.get("reason", ""),
                      det.get("note", "")]
            for col, v in enumerate(values, start=1):
                style_body_cell(ws.cell(row=row, column=col, value=v), row)
            row += 1
        for der in item.get("derived_details", []):
            if der.get("deduplicated"):
                continue
            values = [item.get("element", ""), item.get("designation", ""), der.get("target_sheet", ""),
                      der.get("target_room", ""), "derived", "", der.get("derived_id", ""), "derived from a type definition",
                      f"from {der.get('source_detail', '')} ({der.get('derivation_method', '')})"]
            for col, v in enumerate(values, start=1):
                style_body_cell(ws.cell(row=row, column=col, value=v), row)
            row += 1
    if row == first:
        ws.cell(row=row, column=1, value="Nothing to confirm.")


def build_seen_sheet(ws, data):
    """Every tag the drawing shows that AgentCM has no text for: no card, so nothing to accept until it is drawn."""
    ws.title = "Seen, not in AgentCM"
    row = write_title_row(ws, "Seen on the drawing, not in AgentCM: no card to confirm (inside the possible total)",
                          len(SEEN_COLUMNS))
    row = write_note_row(ws, CLEAR_SEEN, len(SEEN_COLUMNS), row)
    apply_header(ws, SEEN_COLUMNS, row=row)
    row += 1
    first = row
    for item, det in seen(data):
        b = det.get("bounding_box") or {}
        box = " ".join(f"{float(b[k]):.3f}" for k in ("x", "y", "width", "height")) if b else ""
        values = [det.get("tag_text", ""), item.get("designation", ""), det.get("sheet", ""), det.get("room", ""), box,
                  det.get("crop", ""), det.get("corrections", ""), det.get("note", "")]
        for col, v in enumerate(values, start=1):
            style_body_cell(ws.cell(row=row, column=col, value=v), row)
        row += 1
    if row == first:
        ws.cell(row=row, column=1, value="Nothing seen that AgentCM lacks.")


def build_rejected_sheet(ws, data):
    """Every tag the PE already turned down in AgentCM: listed once, in no number."""
    ws.title = "Rejected in AgentCM"
    row = write_title_row(ws, "Rejected in AgentCM: the PE turned these down there (not counted)", len(REJECTED_COLUMNS))
    apply_header(ws, REJECTED_COLUMNS, row=row)
    row += 1
    first = row
    for item, det in rejected(data):
        values = [det.get("tag_text", ""), item.get("designation", ""), det.get("sheet", ""), det.get("room", ""),
                  det.get("detection_id", ""), det.get("note", "")]
        for col, v in enumerate(values, start=1):
            style_body_cell(ws.cell(row=row, column=col, value=v), row)
        row += 1
    if row == first:
        ws.cell(row=row, column=1, value="Nothing rejected.")


# ── Main ───────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="QTO Excel Export")
    parser.add_argument("--data", required=True, help="Path to QTO JSON file")
    parser.add_argument("--project", default="", help="Project name for headers")
    parser.add_argument("--scope", default="", help="Scope label (e.g., 'Room Tags')")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    args = parser.parse_args()
    sys.stdout.reconfigure(errors="replace")  # a tag's text never stops the report on a narrow console

    # Load data
    data_path = Path(args.data)
    if not data_path.exists():
        print(f"ERROR: Data file not found: {data_path}", file=sys.stderr)
        sys.exit(1)

    with open(data_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Validate expected schema keys — the skill must produce exact key names
    required_keys = ["line_items", "scope"]
    missing = [k for k in required_keys if k not in data]
    if missing:
        print(f"WARNING: Missing expected keys: {missing}. Output may be empty.", file=sys.stderr)
        print(f"  Found keys: {list(data.keys())}", file=sys.stderr)
        print(f"  See references/qto-output-format.md for required schema.", file=sys.stderr)

    found = problems(data)
    if found:
        print(f"ERROR: {len(found)} problem(s), no workbook written (references/qto-output-format.md):",
              file=sys.stderr)
        for p in found[:20]:
            print(f"  - {p}", file=sys.stderr)
        sys.exit(2)

    project_name = args.project or data.get("project", {}).get("name", "Unknown Project")
    scope_name = args.scope or data.get("scope", {}).get("tag_type", "Tags")

    # Build workbook
    wb = Workbook()

    build_summary_sheet(wb.active, data, project_name, scope_name)
    build_detail_sheet(wb.create_sheet(), data)
    build_typedef_sheet(wb.create_sheet(), data)
    build_completeness_sheet(wb.create_sheet(), data, project_name, scope_name)
    build_to_confirm_sheet(wb.create_sheet(), data)
    build_seen_sheet(wb.create_sheet(), data)
    build_rejected_sheet(wb.create_sheet(), data)

    # Save
    out = safe_output_path(args.output)
    wb.save(str(out))

    # The four numbers, as the skill reports them
    per_item, total = all_counts(data)
    for item, c in per_item:
        print(f"{label(item)}: {numbers(c)}")
    print(f"TOTAL: {numbers(total)}")
    by_group = groups(data)
    if by_group:
        print("By schedule:")
        for g, c, codes in by_group:
            listed = f" ({', '.join(codes)})" if g == NOT_IN_SCHEDULE else ""
            print(f"  {g}: {numbers(c)}{listed}")
    print(caveat(data))
    on_drawing = seen(data)
    if on_drawing:
        print(f"Seen on the drawing, not in AgentCM: {len(on_drawing)}")
        for item, det in on_drawing:
            print(f"  {det.get('tag_text') or label(item)} on {det.get('sheet', '?')}: {det.get('crop', '')}"
                  + (f" ({det['corrections']})" if det.get("corrections") else ""))
        print(f"  {CLEAR_SEEN}")
    turned_down = rejected(data)
    if turned_down:
        print(f"Rejected in AgentCM, not counted: {len(turned_down)}")
        for item, det in turned_down:
            print(f"  {det.get('tag_text') or label(item)} on {det.get('sheet', '?')} ({det.get('detection_id', '')})")
    print(f"OK: {out}")


if __name__ == "__main__":
    main()
