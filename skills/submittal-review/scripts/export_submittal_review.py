#!/usr/bin/env python3
"""Export a submittal review to Excel, compute the suggested GC disposition, and write
markup items for annotate_pdf.py.

    export_submittal_review.py --review-dir <dir> --output "<folder>/<no> R0 - GC Review.xlsx"
                               [--annotations <dir>/markup] [--allow-gaps]

Refuses to export while check_review.py reports gaps, unless --allow-gaps (the
workbook is then marked INCOMPLETE). Never overwrites: existing outputs get _v2, _v3.
Rules for the disposition: ../references/review-writing.md
"""

import argparse
import json
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_review import check, load_batches, load_json, project_root as find_project_root  # noqa: E402
from shared import safe_output_path  # noqa: E402

SEV_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}
SEV_FILL = {"critical": "F4B6B6", "high": "F8D7A8", "medium": "FBEFB5", "low": "E3EEF7"}
STATUS_FILL = {"finding": "F8D7A8", "unverifiable": "FBEFB5", "na": "EEEEEE", "pass": "D9EFD9",
               "read": "D9EFD9", "unavailable": "FBEFB5"}


def suggest_disposition(findings, inventory, compliance, report):
    sub = [f for f in findings if f.get("owner") == "subcontractor"]
    missing = [e for e in inventory.get("elements", []) if e.get("status") == "missing"]
    missing += [p for p in inventory.get("package_items", []) if p.get("status") == "missing"]
    if missing:
        action = "Return to subcontractor — incomplete; revise and resubmit"
        why = f"{len(missing)} required element(s) or Part 1 item(s) not submitted"
    elif any(f.get("severity") in ("critical", "high") for f in sub):
        action = "Return to subcontractor — revise and resubmit"
        n = sum(1 for f in sub if f.get("severity") in ("critical", "high"))
        why = f"{n} critical/high finding(s) for the subcontractor to correct"
    elif sub:
        action = "Forward to design team with GC comments"
        why = f"{len(sub)} medium/low finding(s) noted for correction"
    else:
        action = "Forward to design team — no GC exceptions"
        why = "No subcontractor findings"
    extras = []
    design = [f for f in findings if f.get("owner") == "design_team"]
    if design:
        extras.append(f"plus {len(design)} RFI candidate(s) for the design team")
    open_serious = [r for r in compliance if r.get("status") == "open" and r.get("severity") in ("critical", "high")]
    if open_serious:
        extras.append(f"{len(open_serious)} high/critical code question(s) open: research before release for fabrication")
    unver = report.get("unverifiable") or []
    if unver:
        extras.append(f"{len(unver)} item(s) could not be verified")
    return action, why, extras


def _style_header(ws, row, ncols):
    from openpyxl.styles import Font, PatternFill
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3A4A5C")


def _table(ws, headers, rows, widths, fill_col=None, fills=None):
    from openpyxl.styles import Alignment, PatternFill
    ws.append(headers)
    _style_header(ws, ws.max_row, len(headers))
    start = ws.max_row + 1
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=start, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill_col is not None and fills:
            key = str(row[fill_col].value or "").lower()
            if key in fills:
                row[fill_col].fill = PatternFill("solid", fgColor=fills[key])
    ws.freeze_panes = ws.cell(row=start, column=1)


def build_workbook(d, state, findings, records, package, inventory, routing, compliance, prior, report, disposition):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    files = (state.get("submittal") or {}).get("files") or []
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    sub = state.get("submittal") or {}
    action, why, extras = disposition
    s = report["stats"]
    lines = [
        ("GC SUBMITTAL REVIEW — DRAFT FOR PE DECISION", None),
        ("Submittal", f"{sub.get('number', '')} R{sub.get('revision', 0)} — {sub.get('title', '')}"),
        ("Subcontractor", sub.get("subcontractor", "")),
        ("Spec section(s)", ", ".join(state.get("sections", []))),
        ("Submittal types", ", ".join(sub.get("types", []))),
        ("Facility types", ", ".join(state.get("facility_types", [])) or "none"),
        ("Suggested GC action", action),
        ("Because", why),
        ("Also", "; ".join(extras) or "—"),
        ("Findings", ", ".join(f"{sev}: {sum(1 for f in findings if f.get('severity') == sev)}"
                               for sev in SEV_ORDER)),
        ("By owner", ", ".join(f"{o}: {sum(1 for f in findings if f.get('owner') == o)}"
                               for o in ("subcontractor", "gc", "design_team", "owner"))),
        ("Elements", f"{s.get('elements_reviewed')} of {s.get('elements_submitted')} submitted reviewed; "
                     f"{s.get('elements_missing')} missing; {s.get('elements_in_trace')} in the contract documents"),
        ("Governing references", f"{s.get('refs_read')} read, {s.get('refs_unavailable')} not in the set, "
                                 f"of {s.get('refs_governing')}"),
        ("Could not verify", f"{s.get('unverifiable')} item(s) — see the Unverifiable sheet"),
        ("Code questions", f"{s.get('open_questions')} open, {s.get('researched_questions')} researched"),
        ("Review gate", "PASSED" if report["ok"] else f"INCOMPLETE — {len(report['gaps'])} gaps"),
    ]
    for note in state.get("notes", []) or []:
        lines.append(("Assumption", note))
    for k, v in lines:
        ws.append([k, v])
    ws["A1"].font = Font(bold=True, size=13)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110

    def page_ref(f):
        sp = f.get("submittal") or {}
        if not isinstance(sp.get("page"), int):
            return ""
        name = Path(files[sp.get("file", 0)]).name if files and sp.get("file", 0) < len(files) else "file"
        return f"{name} p{sp['page']}"

    ws = wb.create_sheet("Findings")
    rows = []
    for f in sorted(findings, key=lambda f: (SEV_ORDER.get(f.get("severity"), 9), f.get("id", ""))):
        req = f.get("requirement") or {}
        rows.append([f.get("id"), f.get("severity"), f.get("owner"), f.get("kind"), f.get("grade"),
                     f.get("element"), f.get("finding"), f"{req.get('source', '')}: {req.get('text', '')}".strip(": "),
                     page_ref(f), f.get("action"), "yes" if f.get("rfi_candidate") else ""])
    _table(ws, ["ID", "Severity", "Owner", "Kind", "Grade", "Element", "Finding", "Governing requirement",
                "Submittal", "Action", "RFI candidate"],
           rows, [9, 10, 14, 15, 13, 14, 60, 50, 22, 50, 10], fill_col=1, fills=SEV_FILL)

    ws = wb.create_sheet("Routing")
    rows = [[r.get("trade"), ", ".join(r.get("sections", []) or []), r.get("send"), r.get("ask"),
             r.get("confirm", ""), r.get("gate", ""), r.get("need_by") or "", r.get("schedule_activity", ""),
             ", ".join(r.get("findings", []) or [])]
            for r in sorted(routing, key=lambda r: (r.get("need_by") or "9999", r.get("trade") or ""))]
    _table(ws, ["Trade", "Sections", "Send them", "Ask", "Confirm who", "Gate", "Need by", "Schedule activity",
                "Findings"], rows, [22, 18, 40, 50, 40, 20, 12, 28, 16])

    ws = wb.create_sheet("Code questions")
    rows = [[r.get("severity"), r.get("status"), r.get("question"), r.get("authority", ""),
             ", ".join(r.get("affects", []) or []), r.get("research") or r.get("source") or "",
             r.get("result") or "", r.get("finding") or ""] for r in compliance]
    _table(ws, ["Severity", "Status", "Question", "Authority", "Affects", "Research / source", "Result", "Finding"],
           rows, [10, 12, 80, 26, 18, 50, 12, 10], fill_col=1, fills={"open": "FBEFB5", "researched": "D9EFD9"})

    ws = wb.create_sheet("Elements")
    inv = {e.get("id"): e for e in inventory.get("elements", [])}
    rec = {r.get("element"): r for r in records}
    trace = load_json(d / "trace.json", {})
    rows = []
    for e in trace.get("elements", []):
        i = inv.get(e.get("id"), {})
        r = rec.get(e.get("id"), {})
        pages = ", ".join(f"p{p.get('page')}" for p in i.get("pages", []) or [])
        read = sum(1 for x in r.get("refs", []) or [] if x.get("status") == "read")
        rows.append([e.get("id"), e.get("type"), e.get("label"), "; ".join(e.get("cd_refs", []) or []),
                     i.get("status", "not inventoried"), pages,
                     f"{read}/{len(e.get('cd_refs', []) or [])}" if r else "", ", ".join(r.get("findings", []) or [])])
    for x in inventory.get("extra_items", []):
        pages = ", ".join(f"p{p.get('page')}" for p in x.get("pages", []) or [])
        rows.append(["(not in CDs)", "", x.get("label"), "", "extra", pages, "", x.get("note", "")])
    _table(ws, ["Element", "Type", "Description", "Contract document references", "Submittal", "Pages",
                "References read", "Findings"], rows, [16, 20, 40, 70, 14, 20, 12, 20])

    ws = wb.create_sheet("Package")
    rows = [[r.get("item"), r.get("status"), r.get("finding") or "", r.get("note", "")] for r in package]
    _table(ws, ["Question or reconciliation", "Status", "Finding", "Note"], rows, [50, 13, 10, 80],
           fill_col=1, fills=STATUS_FILL)

    ws = wb.create_sheet("Unverifiable")
    rows = [[el, u.get("what"), u.get("missing")] for el, u in report.get("unverifiable") or []]
    rows += [[r.get("element"), x.get("ref"), x.get("note", "")] for r in records for x in r.get("refs", []) or []
             if x.get("status") == "unavailable"]
    _table(ws, ["Element", "Could not verify", "What is missing"], rows, [16, 60, 70])

    if prior:
        ws = wb.create_sheet("Prior comments")
        rows = [[p.get("prior_finding"), p.get("status"), p.get("evidence", ""), p.get("finding") or ""] for p in prior]
        _table(ws, ["Prior finding", "Status", "Evidence", "New finding"], rows, [14, 16, 80, 12])

    if not report["ok"]:
        ws = wb.create_sheet("Gate gaps")
        _table(ws, ["Gap"], [[g] for g in report["gaps"]], [140])
    return wb


def markup_items(findings, files, project_root):
    """One JSON list per submittal file, in annotate_pdf.py's item format."""
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            fitz = None
    per_file = {}
    stack = {}
    for f in sorted(findings, key=lambda f: (SEV_ORDER.get(f.get("severity"), 9), f.get("id", ""))):
        sp = f.get("submittal") or {}
        page = (f.get("markup") or {}).get("page") or sp.get("page")
        if not isinstance(page, int):
            continue
        idx = sp.get("file", 0)
        if idx >= len(files):
            continue
        color = "red" if f.get("severity") in ("critical", "high") else "orange"
        text = f"{f.get('id')} [{f.get('severity')}] {f.get('finding')} ACTION: {f.get('action')}"
        items = per_file.setdefault(idx, [])
        rect = (f.get("markup") or {}).get("rect")
        if rect:
            items.append({"page": page, "shape": "cloud", "rect": rect, "color": color,
                          "label": f.get("id"), "content": text, "subject": "GC review (draft)"})
        width, height = 612.0, 792.0
        if fitz is not None:
            path = Path(files[idx])
            path = path if path.is_absolute() else project_root / path
            if path.exists():
                with fitz.open(path) as doc:
                    if 0 < page <= len(doc):
                        width, height = doc[page - 1].rect.width, doc[page - 1].rect.height
        n = stack.get((idx, page), 0)
        stack[(idx, page)] = n + 1
        box_w = min(260.0, width * 0.32)
        x1 = width - 18
        y0 = 18 + n * 58
        if y0 + 54 > height - 18:
            y0 = height - 72
        items.append({"page": page, "shape": "text", "rect": [x1 - box_w, y0, x1, y0 + 54], "color": color,
                      "fontsize": 7, "label": text[:400], "content": text, "subject": "GC review (draft)"})
    return per_file


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--review-dir", required=True)
    ap.add_argument("--output", required=True, help="Excel file to write (never overwritten)")
    ap.add_argument("--annotations", help="Folder for markup item files, one per submittal file")
    ap.add_argument("--allow-gaps", action="store_true", help="Export even though the review gate reports gaps")
    args = ap.parse_args()

    d = Path(args.review_dir).resolve()
    report = check(d)
    if not report["ok"] and not args.allow_gaps:
        print(f"Review gate failed ({len(report['gaps'])} gaps). Run check_review.py and fix them first.")
        for g in report["gaps"][:20]:
            print(f"GAP   {g}")
        return 1

    state = yaml.safe_load((d / "state.yaml").read_text(encoding="utf-8")) or {}
    findings = load_batches(d, "findings")
    records = load_batches(d, "review")
    package = load_json(d / "package_review.json", [])
    inventory = load_json(d / "inventory.json", {})
    routing = load_json(d / "routing.json", [])
    compliance = load_json(d / "compliance.json", [])
    prior = load_json(d / "prior.json", [])

    disposition = suggest_disposition(findings, inventory, compliance, report)
    wb = build_workbook(d, state, findings, records, package, inventory, routing, compliance, prior, report, disposition)
    out = safe_output_path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")

    if args.annotations:
        project_root = find_project_root(d)
        files = (state.get("submittal") or {}).get("files") or []
        folder = Path(args.annotations)
        folder.mkdir(parents=True, exist_ok=True)
        for idx, items in markup_items(findings, files, project_root).items():
            path = folder / f"{Path(files[idx]).stem}.json"
            path.write_text(json.dumps(items, indent=2), encoding="utf-8")
            print(f"Wrote {path} ({len(items)} markups)")

    action, why, extras = disposition
    print(f"Suggested GC action: {action} — {why}" + (f"; {'; '.join(extras)}" if extras else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
