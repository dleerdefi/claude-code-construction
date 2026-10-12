#!/usr/bin/env python3
"""Write evals/plugin/agentcm-mode/ (AgentCM's contract fixture and its scripts) and the agentcm-mode-<skill> cases
from the fixture AgentCM's scripts/writeContractFixture.ts writes. Run it whenever that fixture changes:

    bin/construction-python dev/make_agentcm_cases.py <fixture root: the folder holding .construction/>

The facts the graders check (sheet counts, a door's size, the accepted and pending tags) are read from the
fixture here, so a new fixture only needs a re-run. Review the diff before committing."""
import json, re, shutil, sys, textwrap
from pathlib import Path

W = Path(__file__).resolve().parents[1]; FIX = Path(sys.argv[1])
P = W / "evals" / "plugin"
A = P / "agentcm-mode"
STUB = "http://127.0.0.1:47817/"


def write(rel, text):
    p = P / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(textwrap.dedent(text).lstrip("\n"), encoding="utf-8", newline="\n")


# ---- the fixture, copied as it is -----------------------------------------------------------------
for sub in ("construction", "project"):
    if (A / sub).exists():
        shutil.rmtree(A / sub)
for name in ("FIXTURE_SOURCE", "api-responses.json"):
    if (A / name).exists():
        (A / name).unlink()
shutil.copytree(FIX / ".construction", A / "construction")
for entry in sorted(FIX.iterdir()):
    if entry.name == ".construction":
        continue
    if entry.name in ("FIXTURE_SOURCE", "api-responses.json"):
        shutil.copy2(entry, A / entry.name)
    elif entry.is_dir():
        shutil.copytree(entry, A / "project" / entry.name)
    else:
        (A / "project").mkdir(parents=True, exist_ok=True)
        shutil.copy2(entry, A / "project" / entry.name)

C = A / "construction"
INDEX = (C / "INDEX.md").read_text(encoding="utf-8")


def heading_count(name):
    m = re.search(rf"^## {re.escape(name)} \((\d+)\)", INDEX, re.M)
    return int(m.group(1)) if m else 0


def table(heading):
    part = INDEX.split(f"## {heading}", 1)[1].split("\n## ", 1)[0] if f"## {heading}" in INDEX else ""
    rows = [r for r in part.splitlines() if r.startswith("| ") and not set(r) <= set("|- ")]
    return [[c.strip() for c in r.strip("|").split("|")] for r in rows[1:]]


SHEETS = [r[0] for r in table("Drawing sheets")]
SCHEDULES = table("Schedules and legends")
SECTIONS = table("Specification sections")
UNREAD = table("Spec manuals not read")
READ_MANUALS = sorted({r[4] for r in SECTIONS})
DRAWINGS = sorted(f"drawings/{f.name}" for f in (A / "project" / "drawings").glob("*.pdf"))
ROOMS = re.search(r"^Rooms known: (\d+)", INDEX, re.M)
door = next((l.split("\t") for l in (C / "schedules" / "A-601-1.tsv").read_text(encoding="utf-8").splitlines()
             if "\tD103\t" in l), None)
D103_SIZE = door[2] if door else None
XREF_A102 = [l.split("\t") for l in (C / "xrefs.tsv").read_text(encoding="utf-8").splitlines()[1:] if l.split("\t")[5:6] == ["A-102"]]
RESP = json.loads((A / "api-responses.json").read_text(encoding="utf-8")) if (A / "api-responses.json").exists() else {}


def tsv(path):
    rows = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    head = rows[0].split("\t") if rows else []
    return [dict(zip(head, r.split("\t"))) for r in rows[1:]]


def box(row):
    return tuple(float(row[k]) for k in ("x", "y", "w", "h"))


def near(a, b, pad=0.01):
    return a[0] - pad <= b[0] + b[2] and b[0] - pad <= a[0] + a[2] and a[1] - pad <= b[1] + b[3] and b[1] - pad <= a[1] + a[3]


def room_tags(sheets):
    """(accepted, pending) room tags on `sheets`, from AgentCM's files as tag-audit reads them: accepted, the
    .groups.tsv room_tag lines and the accepted room lines of tags.tsv outside every such box; pending, the .tsv
    room_tag lines marked pending, neighbouring lines one tag."""
    accepted = pending = 0
    points = [t for t in tsv(C / "tags.tsv") if t.get("kind") == "room" and t.get("how") == "accepted"]
    for s in sheets:
        groups = [box(g) for g in tsv(C / "sheets" / f"{s}.groups.tsv") if g.get("type") == "room_tag"]
        accepted += len(groups) + sum(1 for t in points if t.get("sheet") == s and not any(
            g[0] <= float(t["x"]) <= g[0] + g[2] and g[1] <= float(t["y"]) <= g[1] + g[3] for g in groups))
        lines = [box(l) for l in tsv(C / "sheets" / f"{s}.tsv") if l.get("tag") == "room_tag" and l.get("how") == "pending"]
        seen = set()
        for i in range(len(lines)):
            if i in seen:
                continue
            pending += 1
            stack = [i]
            while stack:
                j = stack.pop()
                if j not in seen:
                    seen.add(j)
                    stack += [k for k in range(len(lines)) if k not in seen and near(lines[j], lines[k])]
    return accepted, pending


ROOM_ACCEPTED, ROOM_PENDING = room_tags(["A-101", "A-102"])

# ---- fixture scripts ------------------------------------------------------------------------------
write("agentcm-mode/README.md", f"""
    # AgentCM contract fixture

    A project folder as AgentCM leaves it, written by AgentCM's own `scripts/writeContractFixture.ts`
    (`FIXTURE_SOURCE` gives the AgentCM commit and the contract version), and what the agentcm-mode cases
    need to run against it:

    | Here | In the workspace | What it is |
    |---|---|---|
    | `construction/` | `.construction/` | AgentCM's files (stored without the dot: the repository ignores `.construction/`) |
    | `project/` | the project folder | The PDFs INDEX.md names (drawings, specification manuals) |
    | `api-responses.json` | | AgentCM's recorded answers, which `stub_api.py` serves at the fixture's `api_url` |
    | `setup-fixture.sh` | | Each case's setup: copies the fixture, makes `query.sh` fail, records a baseline, starts or stops the stub |
    | `stub_api.py`, `teardown.sh` | | The stand-in for AgentCM's API, and its stop |
    | `baseline.py` | `.eval-api/fixture-baseline.json` | The fixture's files as setup left them, for the writes check |

    **AgentCM closed by default.** Skills read AgentCM's files and never its database, and a read-only skill
    must work with AgentCM closed. So every case runs with nothing answering at `api_url` and `query.sh`
    failing, except the cases that send something back: those start the stub (`setup-fixture.sh --api`)
    and list `{STUB}` in `allow_urls`, the one address curl may reach. A writing skill's offline case keeps
    the stub off and expects the payload in `.construction/skills/<skill>/outbox/`.

    **The stub** answers GETs from `api-responses.json` (contract 1 promises one read, a schedule's rows),
    keeps and answers the three write endpoints, and answers nothing else. A reconcile is answered from
    the schedule's stored rows (the recorded rows read), never from another request's recorded answer.
    A tag ingest is answered by
    each detection's box over the workspace's sheet files, as D-249 has AgentCM do: `created`,
    `already_held` (with the holding card), `no_free_text`, or `rejected` for a tag the recorded ingest
    shows rejected. Every request, with a write's answer, goes to `.eval-api/requests.jsonl`, which the
    trace graders read (the right endpoint, the contract's payload, the sheet id from
    `index/sheet_index.yaml`, what the skill was told).

    **Waiting cases.** A case whose `setup.sh` names a need the fixture does not meet (`specs`,
    `unread-manual`, `groups`, `pending`, `pdfs`, `api`, `api-url`, `schedule-id`, `flooring-specs`,
    `submittal`) is skipped with "waiting for the contract fixture"; nothing is faked. The 0.6.0 gate runs
    `--tag agentcm-mode --no-skip`, so a waiting case fails it.

    **To update:** `bin/construction-python dev/make_agentcm_cases.py <AgentCM's fixture folder>`; it copies `.construction/`,
    the PDFs, `api-responses.json` and `FIXTURE_SOURCE`, and rewrites the cases' expected facts from it.
""")

write("agentcm-mode/setup-fixture.sh", r"""
    #!/usr/bin/env bash
    # Copy AgentCM's contract fixture into the eval workspace, with AgentCM closed unless --api.
    # Usage: setup-fixture.sh [--api] [needs...]
    #   needs: specs unread-manual groups pending pdfs api api-url schedule-id flooring-specs submittal
    # A need the fixture does not meet exits 3: the harness reports the case as skipped.
    set -euo pipefail
    HERE="$(cd "$(dirname "${BASH_SOURCE[0]}")" && pwd)"
    PY="$HERE/../../../bin/construction-python"
    SRC="$HERE/construction"
    start_api=0
    if [ "${1:-}" = "--api" ]; then start_api=1; shift; set -- api "$@"; fi
    missing=()
    for need in "$@"; do
      case "$need" in
        specs) ls "$SRC"/specs/*.txt >/dev/null 2>&1 && grep -q '^## Specification sections' "$SRC/INDEX.md" || missing+=("spec sections") ;;
        unread-manual) grep -q '^## Spec manuals not read' "$SRC/INDEX.md" || missing+=("a manual AgentCM could not read") ;;
        groups) ls "$SRC"/sheets/*.groups.tsv >/dev/null 2>&1 || missing+=("accepted groups") ;;
        pending) grep -q "$(printf '\tpending\t')" "$SRC"/sheets/*.tsv 2>/dev/null || missing+=("a pending group") ;;
        pdfs) ls "$HERE"/project/drawings/*.pdf >/dev/null 2>&1 || missing+=("the sheets' PDFs") ;;
        api) [ -f "$HERE/api-responses.json" ] && grep -q '127.0.0.1:47817' "$SRC/database.yaml" || missing+=("AgentCM's recorded answers at 127.0.0.1:47817") ;;
        api-url) grep -q '127.0.0.1:47817' "$SRC/database.yaml" || missing+=("an api_url at 127.0.0.1:47817") ;;
        schedule-id) grep -q 'schedule_id=[0-9a-f]\{8\}-' "$SRC"/schedules/*.tsv 2>/dev/null || missing+=("an ingested schedule (schedule_id=)") ;;
        flooring-specs) grep -q '^| 09 6' "$SRC/INDEX.md" || missing+=("flooring sections matching the BP-09 bids") ;;
        submittal) [ -d "$HERE/project/06 - Submittals" ] || missing+=("a submittal against one of its sections") ;;
        *) echo "setup-fixture.sh: unknown need '$need'" >&2; exit 2 ;;
      esac
    done
    if [ "${#missing[@]}" -gt 0 ]; then
      printf 'waiting for the contract fixture: it has no %s\n' "$(IFS=';'; echo "${missing[*]}")" >&2
      exit 3
    fi
    mkdir -p .construction
    cp -R "$SRC/." .construction/
    if [ -d "$HERE/project" ]; then cp -R "$HERE/project/." .; fi
    # AgentCM closed: its database does not answer (skills never query it anyway).
    printf '#!/bin/sh\necho "AgentCM is not running: its database does not answer" >&2\nexit 1\n' > .construction/query.sh
    printf 'Write-Error "AgentCM is not running: its database does not answer"\nexit 1\n' > .construction/query.ps1
    if [ "$start_api" = 1 ]; then
      "$PY" "$HERE/stub_api.py" start --responses "$HERE/api-responses.json" --workspace .
    else
      "$PY" "$HERE/stub_api.py" stop >/dev/null || true
    fi
""")

write("agentcm-mode/baseline.sh", r"""
    #!/usr/bin/env bash
    # The last step of an agentcm-mode case's setup: record the fixture as setup left it.
    set -euo pipefail
    HERE="$(cd "$(dirname "${BASH_SOURCE[0]}")" && pwd)"
    "$HERE/../../../bin/construction-python" "$HERE/baseline.py" .
""")

write("agentcm-mode/baseline.py", '''
    #!/usr/bin/env python3
    """Record every file under .construction/ outside skills/ and crops/ with its hash, as setup left it, in
    .eval-api/fixture-baseline.json: the writes check compares the run's end state against it."""
    import hashlib
    import json
    import sys
    from pathlib import Path

    ws = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    base = ws / ".construction"
    files = {}
    for p in sorted(base.rglob("*")):
        rel = p.relative_to(base).as_posix()
        if p.is_file() and not rel.startswith(("skills/", "crops/")):
            files[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
    out = ws / ".eval-api" / "fixture-baseline.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(files, indent=1), encoding="utf-8")
    print(f"baseline: {len(files)} file(s) under .construction/")
''')

write("agentcm-mode/teardown.sh", r"""
    #!/usr/bin/env bash
    # After an agentcm-mode run: stop the stand-in for AgentCM's API, if one is answering.
    HERE="$(cd "$(dirname "${BASH_SOURCE[0]}")" && pwd)"
    "$HERE/../../../bin/construction-python" "$HERE/stub_api.py" stop
""")

write("agentcm-mode/make_edited_schedule.py", '''
    #!/usr/bin/env python3
    """For the reconcile case: the workbook a PE edits after exporting AgentCM's ingested schedule.

    The ingested schedule is the one whose schedules/<sheet>-<n>.tsv comment line carries schedule_id=<uuid>;
    its rows are AgentCM's recorded answer to GET .../schedules/<id>/rows. The workbook is written by the
    plugin's own schedule_to_xlsx.py (the anchored export: _agentcm_meta, _row_key), then one cell is edited:
    the first row's last column other than the row key gets " (REV)" appended. expectation() gives that edit
    without writing anything, for the check.

    Usage: make_edited_schedule.py <workspace>   writes "<sheet> Schedule (PE edits).xlsx" there
    """
    import json
    import re
    import sys
    import tempfile
    from pathlib import Path

    HERE = Path(__file__).resolve().parent
    SCRIPTS = HERE.parents[2] / "skills" / "schedule-extractor" / "scripts"


    def ingested():
        """(sheet, schedule_id, rows) of the fixture's ingested schedule."""
        for f in sorted((HERE / "construction" / "schedules").glob("*.tsv")):
            first = f.read_text(encoding="utf-8").splitlines()[0]
            m = re.search(r"schedule_id=([0-9a-f-]{36})", first)
            if m:
                responses = json.loads((HERE / "api-responses.json").read_text(encoding="utf-8"))
                rows = next(v for k, v in responses.items() if k.startswith("GET ") and k.endswith(f"/schedules/{m.group(1)}/rows"))
                return f.stem.rsplit("-", 1)[0], m.group(1), rows
        raise SystemExit("no ingested schedule (schedule_id=<uuid>) in the fixture")


    def key_column(rows):
        cols = list(rows[0]["data"].keys())
        return next((c for c in cols if all(str(r["data"].get(c)) == str(r["entityIdentifier"]) for r in rows)), None)


    def expectation():
        sheet, sid, rows = ingested()
        key = key_column(rows)
        col = [c for c in rows[0]["data"] if c != key][-1]
        return {"sheet": sheet, "schedule_id": sid, "row_key": str(rows[0]["entityIdentifier"]), "column": col,
                "value": f"{rows[0]['data'].get(col) or ''} (REV)".strip()}


    def main(ws):
        from openpyxl import load_workbook
        sys.path.insert(0, str(SCRIPTS))
        from schedule_to_xlsx import schedule_to_xlsx
        sheet, sid, rows = ingested()
        exp = expectation()
        sheet_ids = (HERE / "construction" / "index" / "sheet_index.yaml").read_text(encoding="utf-8")
        sheet_id = re.search(r'id: "([^"]+)"\\s+sheet_number: "' + re.escape(sheet) + '"', sheet_ids).group(1)
        out = Path(ws) / f"{sheet} Schedule (PE edits).xlsx"
        with tempfile.TemporaryDirectory() as d:
            data = Path(d) / "rows.json"
            data.write_text(json.dumps([{**{k: v for k, v in r["data"].items()}} for r in rows]), encoding="utf-8")
            schedule_to_xlsx(str(data), output=str(out), schedule_type="schedule", sheet=sheet, schedule_id=sid,
                             sheet_id=sheet_id, row_key_column=key_column(rows) or "")
        wb = load_workbook(out)
        ws_data = next(wb[n] for n in wb.sheetnames if n != "_agentcm_meta")
        header_row = next(r for r in range(1, ws_data.max_row + 1) if ws_data.cell(r, 1).value == "_row_key")
        headers = {ws_data.cell(header_row, c).value: c for c in range(2, ws_data.max_column + 1)}
        row = next(r for r in range(header_row + 1, ws_data.max_row + 1) if str(ws_data.cell(r, 1).value) == exp["row_key"])
        ws_data.cell(row, headers[exp["column"]], exp["value"])
        wb.save(out)
        print(f"edited {out.name}: {exp['row_key']} / {exp['column']} = {exp['value']!r}")


    if __name__ == "__main__":
        main(sys.argv[1] if len(sys.argv) > 1 else ".")
''')

# ---- shared checks --------------------------------------------------------------------------------
write("_checks/agentcm_mode_writes.py", '''
    #!/usr/bin/env python3
    """AgentCM mode: nothing written under .construction/ outside skills/ and crops/.

    Setup records the fixture as it left it (.eval-api/fixture-baseline.json); a file outside skills/ and
    crops/ that is new or differs from it fails. On a regrade the snapshot holds only files the run wrote,
    so any such file there fails. crops/ is crop.sh's; that nothing else writes there is graded from the
    trace (crops-through-crop-sh).
    """
    import hashlib
    import json
    import os
    import sys
    from pathlib import Path

    workspace = Path(os.environ.get("EVAL_WORKSPACE", "."))
    base = workspace / ".construction"
    baseline_file = workspace / ".eval-api" / "fixture-baseline.json"
    baseline = json.loads(baseline_file.read_text(encoding="utf-8")) if baseline_file.is_file() else {}
    bad = []
    if base.is_dir():
        for p in sorted(base.rglob("*")):
            rel = p.relative_to(base).as_posix()
            if not p.is_file() or rel.startswith(("skills/", "crops/")):
                continue
            if baseline.get(rel) != hashlib.sha256(p.read_bytes()).hexdigest():
                bad.append(f".construction/{rel}")
    if bad:
        print(f"FAIL: {len(bad)} file(s) under .construction/ written outside skills/ and crops/: {bad[:10]}")
        sys.exit(1)
    print("PASS: nothing under .construction/ written outside skills/ and crops/")
''')

write("_checks/agentcm_trace.py", '''
    """Trace and request helpers for the agentcm-mode cases: what a stand-in cannot prove from results is graded
    from the calls made.

    Each case's checks/*.py imports this (sys.path to evals/plugin/_checks) and calls one check. EVAL_TRACE is
    the run's trace.jsonl, EVAL_WORKSPACE its workspace (or regrade snapshot), whose .eval-api/requests.jsonl
    the stub wrote; the fixture's own files are read from evals/plugin/agentcm-mode/construction/.
    """
    import json
    import os
    import re
    import sys
    from pathlib import Path

    import yaml

    HERE = Path(__file__).resolve().parent
    FIXTURE = HERE.parent / "agentcm-mode" / "construction"
    WORKSPACE = Path(os.environ.get("EVAL_WORKSPACE", "."))
    NUM = r"(0?\\.\\d+|0|1(?:\\.0+)?)"
    CROP = re.compile(r"crop\\.(?:sh|ps1)[\\"']?\\s+[\\"']?([A-Za-z0-9._-]+)[\\"']?\\s+" + r"\\s+".join([NUM] * 4))
    READS = ("/rows",)  # contract 1's one promised read
    CANONICAL = {"room_tag", "door_tag", "element_tag", "simple_callout", "multi_detail_callout", "grid_label",
                 "view_label", "note_block", "schedule_block", "schedule_zone", "annotation", "viewport_highlight"}
    WRITES = ("schedules/ingest", "schedules/reconcile", "tag-detections/ingest")
    BOX = ("x", "y", "width", "height")


    def finish(ok, passed, failed):
        print(("PASS: " + passed) if ok else ("FAIL: " + failed))
        sys.exit(0 if ok else 1)


    def bash_commands():
        out = []
        for line in open(os.environ["EVAL_TRACE"], encoding="utf-8"):
            r = json.loads(line)
            if r.get("type") == "tool_use" and r.get("name") in ("Bash", "PowerShell"):
                out.append(r["input"].get("command", ""))
        return out


    def requests():
        """What reached the stand-in API: [{method, path, query, body, status}], in order."""
        p = WORKSPACE / ".eval-api" / "requests.jsonl"
        return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()] if p.is_file() else []


    def posts(endpoint):
        return [r for r in requests() if r["method"] == "POST" and r["path"].endswith("/" + endpoint)]


    def sheet_ids():
        """{sheet_number: id} from the fixture's index/sheet_index.yaml."""
        data = yaml.safe_load((FIXTURE / "index" / "sheet_index.yaml").read_text(encoding="utf-8")) or {}
        return {str(s.get("sheet_number")): str(s.get("id")) for s in data.get("sheets") or [] if s.get("id")}


    def project_id():
        return str((yaml.safe_load((FIXTURE / "database.yaml").read_text(encoding="utf-8")) or {}).get("project_id"))


    # -- crop.sh ------------------------------------------------------------------------------------
    def crops():
        """(sheet, x, y, w, h) for every crop.sh / crop.ps1 call."""
        return [(m.group(1), *(float(v) for v in m.groups()[1:])) for cmd in bash_commands() for m in CROP.finditer(cmd)]


    def sheet_lines(sheet):
        """(x, y, w, h, text) for each line of the fixture's sheets/<sheet>.tsv."""
        p = FIXTURE / "sheets" / f"{sheet}.tsv"
        out = []
        for r in (p.read_text(encoding="utf-8").splitlines()[1:] if p.is_file() else []):
            c = r.split("\\t")
            try:
                out.append((float(c[0]), float(c[1]), float(c[2]), float(c[3]), c[-1]))
            except (ValueError, IndexError):
                pass
        return out


    def overlaps(a, b):
        ax, ay, aw, ah = a; bx, by, bw, bh = b
        return ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah


    def check_crop_covers(sheet, text):
        """A crop of `sheet` whose box overlaps the line of the sheet's .tsv holding `text`."""
        targets = [l[:4] for l in sheet_lines(sheet) if l[4].strip() == text]
        cs = crops()
        hits = [c for c in cs if c[0] == sheet and any(overlaps(c[1:], t) for t in targets)]
        finish(bool(hits) and bool(targets), f"crop.sh {hits[0] if hits else ''} covers {text!r} on {sheet}",
               f"no crop.sh call on {sheet} covers {text!r} {targets}; crops made: {cs[:6]}")


    def check_crop_from_sheet_files(sheets):
        """At least one crop of one of `sheets`, on the sheet, overlapping a line its .tsv gives."""
        cs = crops()
        good = [c for c in cs if c[0] in sheets and all(0 <= v <= 1 for v in c[1:]) and c[1] + c[3] <= 1.0001
                and c[2] + c[4] <= 1.0001 and any(overlaps(c[1:], l[:4]) for l in sheet_lines(c[0]))]
        finish(bool(good), f"crop.sh {good[0] if good else ''} uses a box from the sheet files",
               f"no crop.sh call on {sheets} with a box on a line of the sheet files; crops made: {cs[:6]}")


    # -- own extraction -------------------------------------------------------------------------------
    EXTRACTION = re.compile(r"rasterize_page|pdf_text\\.py|split_drawing_set|split_spec_manual|extract_spec_text|find_pages|"
                            r"extract_text_region|pdfplumber|import fitz")


    def check_no_own_extraction(read, drawings):
        """Nothing AgentCM read is extracted again: no extraction command names a manual INDEX.md lists as read
        (`read`) or a drawing PDF (`drawings`, looked at through crop.sh). A manual INDEX.md lists as not read is
        spec-splitter's to split and extract, and so are the section PDFs it splits from it."""
        bad = [(Path(f).name, c[:160]) for c in bash_commands() if EXTRACTION.search(c)
               for f in read + drawings if Path(f).name in c]
        finish(not bad, "nothing AgentCM read is extracted again", f"extraction of what AgentCM read: {bad[:4]}")


    # -- the API --------------------------------------------------------------------------------------
    def check_contract_requests():
        """Every request to the stand-in is one of the contract's reads or writes, for this project, answered."""
        reqs, pid = requests(), project_id()
        bad = [f"{r['method']} {r['path']} ({r['status']})" for r in reqs
               if not r["path"].startswith(f"/api/projects/{pid}/")
               or not any(e in r["path"] for e in (READS if r["method"] == "GET" else WRITES))
               or r["status"] != 200]
        finish(bool(reqs) and not bad, f"{len(reqs)} request(s), each a contract endpoint for project {pid}",
               f"requests outside the contract, another project, or refused: {bad[:6]} (of {len(reqs)})")


    def detection_problems(d):
        out = []
        for k in ("tag_text", "tag_type", "bounding_box", "confidence"):
            if k not in d:
                out.append(f"no {k}")
        b = d.get("bounding_box")
        if not (isinstance(b, dict) and all(isinstance(b.get(k), (int, float)) and 0 <= b[k] <= 1 for k in BOX)):
            out.append(f"bounding_box not {{x, y, width, height}} in 0-1: {b}")
        if "tag_type" in d and d["tag_type"] not in CANONICAL:
            out.append(f"tag_type {d['tag_type']!r} is not one of contract 1's names")
        if d.get("element_ids"):
            out.append("element_ids sent: a detection goes by its box")
        return out


    def tag_payload_problems(body, sheet):
        if not isinstance(body, dict):
            return ["the body is not a JSON object"]
        out = [] if body.get("sheet_id") == sheet_ids().get(sheet) else [f"sheet_id {body.get('sheet_id')}, not {sheet}'s"]
        dets = body.get("detections")
        if not isinstance(dets, list) or not dets:
            return out + ["no detections"]
        return out + [p for d in dets for p in detection_problems(d if isinstance(d, dict) else {})]


    SCHEDULE_KEYS = ("sheetId", "scheduleType", "title", "rowEntityType", "columns", "rows", "sourceMethod", "confidence")
    # the values AgentCM's own schedule types allow (src/shared/scheduleTypes.ts at 04feab4)
    ROW_ENTITY_TYPES = ("room", "door", "window", "fixture", "equipment")
    SOURCE_METHODS = ("pdfplumber", "vision", "hybrid", "manual")


    def schedule_payload_problems(body, sheet):
        if not isinstance(body, dict):
            return ["the body is not a JSON object"]
        out = [f"no {k}" for k in SCHEDULE_KEYS if k not in body]
        if body.get("sheetId") != sheet_ids().get(sheet):
            out.append(f"sheetId {body.get('sheetId')}, not {sheet}'s")
        if "rowEntityType" in body and body["rowEntityType"] not in ROW_ENTITY_TYPES:
            out.append(f"rowEntityType {body['rowEntityType']!r}, not one of AgentCM's {ROW_ENTITY_TYPES}")
        if "sourceMethod" in body and body["sourceMethod"] not in SOURCE_METHODS:
            out.append(f"sourceMethod {body['sourceMethod']!r}, not one of AgentCM's {SOURCE_METHODS}")
        rows = body.get("rows")
        if not isinstance(rows, list) or not rows or not all(isinstance(r, dict) and r.get("entityIdentifier")
                                                             and isinstance(r.get("cells"), dict) for r in rows):
            out.append("rows without entityIdentifier and cells")
        return out


    def check_tag_ingest(sheets):
        """tag-detections/ingest posts, each for one of `sheets` by its id, with the contract's payload."""
        ps = posts("tag-detections/ingest")
        bad = []
        for r in ps:
            sheet = next((s for s in sheets if r["body"] and r["body"].get("sheet_id") == sheet_ids().get(s)), sheets[0])
            bad += tag_payload_problems(r["body"], sheet)
        finish(bool(ps) and not bad, f"{len(ps)} tag ingest(s), sheet ids from sheet_index.yaml, payloads as the contract",
               f"tag ingests: {len(ps)}; problems: {bad[:6]}")


    def check_schedule_ingest(sheet):
        ps = posts("schedules/ingest")
        bad = [p for r in ps for p in schedule_payload_problems(r["body"], sheet)]
        finish(bool(ps) and not bad, f"{len(ps)} schedule ingest(s) for {sheet} by its id, payload as the contract",
               f"schedule ingests: {len(ps)}; problems: {bad[:6]}")


    def check_reconcile(schedule_id, row_key, column, value):
        """One reconcile post for `schedule_id`, carrying sourceFile and the edited cell, that AgentCM took: its
        answer names no error (the stand-in answers from the stored rows, so an error means a row key it lacks)."""
        ps = posts("schedules/reconcile")
        ok = [r for r in ps if isinstance(r["body"], dict) and r["body"].get("scheduleId") == schedule_id
              and r["body"].get("sourceFile") and any(u.get("rowKey") == row_key and u.get("columnKey") == column
                                                      and u.get("newValue") == value for u in r["body"].get("cellUpdates") or [])
              and isinstance(r.get("response"), dict) and not r["response"].get("errors")]
        rows_read = [r for r in requests() if r["method"] == "GET" and r["path"].endswith(f"/schedules/{schedule_id}/rows")]
        finish(bool(ok) and bool(rows_read), f"rows read, then reconcile of {schedule_id} with {row_key}/{column} = {value!r}",
               f"rows reads: {len(rows_read)}; reconciles: {len(ps)}, none with {row_key}/{column} = {value!r} and "
               f"no error: {[(r['body'], r.get('response')) for r in ps][:2]}")


    # -- AgentCM closed -------------------------------------------------------------------------------
    def outbox(skill):
        d = WORKSPACE / ".construction" / "skills" / skill / "outbox"
        out = []
        for p in sorted(d.glob("*.json")) if d.is_dir() else []:
            try:
                out.append((p.name, json.loads(p.read_text(encoding="utf-8"))))
            except json.JSONDecodeError:
                out.append((p.name, None))
        return out


    def check_outbox(skill, kind, sheet):
        """AgentCM closed: nothing reached an API, and the payload waits in the skill's outbox."""
        saved = outbox(skill)
        problems = tag_payload_problems if kind == "tag" else schedule_payload_problems
        good = [n for n, body in saved if body is not None and not problems(body, sheet)]
        reached = requests()
        finish(bool(good) and not reached, f"{good[0] if good else ''} waits in {skill}/outbox/ with the contract's payload",
               f"outbox: {[n for n, _ in saved]}, well-formed for {sheet}: {good}; requests that reached an API: {len(reached)}")


    # -- the takeoff, twice -----------------------------------------------------------------------------
    def answered(endpoint):
        """(sheet id, detection, the stand-in's answer to it) for every detection sent to `endpoint`, in order."""
        out = []
        for r in posts(endpoint):
            body = r["body"] if isinstance(r["body"], dict) else {}
            results = {e.get("index"): e for e in ((r.get("response") or {}).get("results") or []) if isinstance(e, dict)}
            out += [(body.get("sheet_id"), d, results.get(i, {})) for i, d in enumerate(body.get("detections") or [])
                    if isinstance(d, dict)]
        return out


    def check_takeoff_twice(accepted, pending):
        """Two QTO JSONs with the same confirmed, to confirm and possible total per item; confirmed is the
        accepted tags (`accepted`), to confirm holds at least the pending ones. A tag sent again (the files
        do not refresh here) came back already_held, never as a second card; a tag AgentCM answered rejected
        is in each QTO once, as rejected, in no number."""
        sys.path.insert(0, str(HERE.parents[2] / "skills" / "tag-audit-and-takeoff" / "scripts"))
        import qto_to_xlsx as q
        files = sorted((WORKSPACE / ".construction" / "skills" / "tag-audit-and-takeoff" / "qto").glob("*.json"))
        if len(files) < 2:
            finish(False, "", f"{len(files)} QTO JSON(s); a second takeoff was asked for")
        runs, datas = [], []
        for f in files[-2:]:
            data = json.loads(f.read_text(encoding="utf-8"))
            datas.append(data)
            per_item, total = q.all_counts(data)
            runs.append(({q.label(i): c[:1] + c[1:2] + c[3:] for i, c in per_item}, (total[0], total[1], total[3])))
        same = runs[0] == runs[1]
        conf, todo, _ = runs[1][1]

        def one_tag(a, b):
            ba, bb = q.box_of(a[1]), q.box_of(b[1])
            return (a[0] == b[0] and q.norm_text(a[1].get("tag_text")) == q.norm_text(b[1].get("tag_text")) and ba and bb
                    and (q.center_inside(ba, bb) or q.center_inside(bb, ba)))

        sent = answered("tag-detections/ingest")
        second_card = [s[1].get("tag_text") for i, s in enumerate(sent) for later in sent[i + 1:]
                       if one_tag(s, later) and later[2].get("outcome") == "created"]
        # the second run's trace: the files do not refresh here, so it sends a first-run tag again and hears already_held
        created = [i for i, s in enumerate(sent) if s[2].get("outcome") == "created"]
        resent = not created or any(one_tag(sent[i], later) and later[2].get("outcome") == "already_held"
                                    for i in created for later in sent[i + 1:])
        number = {v: k for k, v in sheet_ids().items()}
        turned_down = {(number.get(s), q.norm_text(d.get("tag_text"))) for s, d, a in sent if a.get("outcome") == "rejected"}
        not_listed_once = []
        for k, data in enumerate(datas, start=1):
            for sheet, text in sorted(turned_down):
                found = [d for item in data.get("line_items", []) for d in item.get("instance_details", [])
                         if d.get("sheet") == sheet and q.norm_text(d.get("tag_text")) == text]
                if [q.status_of(d) for d in found] != ["rejected"]:
                    not_listed_once.append(f"run {k}: {text} on {sheet} as {[d.get('status') for d in found]}")
        ok = same and conf == accepted and todo >= pending and resent and not second_card and not not_listed_once
        finish(ok, f"both takeoffs: confirmed {conf}, to confirm {todo}; a tag sent again was already held; "
                   f"rejected listed once: {sorted(t for _, t in turned_down)}",
               f"runs {runs}; expected confirmed {accepted}, to confirm >= {pending}; a first-run tag sent again and "
               f"already held: {resent}; second cards: {second_card[:3]}; rejected not listed once as rejected: "
               f"{not_listed_once[:3]}")
''')

# ---- cases ----------------------------------------------------------------------------------------
UNATTENDED = ("This is an unattended evaluation run in a construction project folder that AgentCM manages. "
              "Nobody can answer questions: make reasonable choices, do not wait for confirmation, and finish the task.")
COMMON = {
    "agentcm-mode": ("regex", 'target: trace\npattern: "\\\\.construction[\\\\\\\\/]+project\\\\.yaml"', 1,
                     "The run read .construction/project.yaml: it decided the mode from AgentCM's marker."),
    "agentcm-files": ("regex", 'target: trace\npattern: "\\\\.construction[\\\\\\\\/]+(INDEX\\\\.md|sheets[\\\\\\\\/]+|schedules[\\\\\\\\/]+|tags\\\\.tsv|xrefs\\\\.tsv|specs[\\\\\\\\/]+|index[\\\\\\\\/]+sheet_index\\\\.yaml)"', 1,
                      "The run read AgentCM's files (INDEX.md, the sheet files, schedules, tags, xrefs, specs, the sheet index)."),
    "no-graph-export": ("regex", 'target: trace\npattern: "graph/export|\\\\.construction[\\\\\\\\/]+rasters|\\\\.construction[\\\\\\\\/]+extractions"\nmatch: not_contains', 1,
                        "Nothing outside the contract: no graph/export, rasters/ or extractions/."),
    "no-db-queries": ("tool_used", 'tool: Bash\ninput_match: "query\\\\.(sh|ps1)|\\\\bpsql\\\\b|PGPASSWORD|query_command"\nmin: 0\nmax: 0', 1,
                      "No database query at all: no query.sh or query.ps1, no psql, password or query_command."),
    "no-stub-control": ("tool_used", 'tool: Bash\ninput_match: "__eval_stub__"\nmin: 0\nmax: 0', 1,
                        "The run never touches the eval stand-in's own controls."),
    "no-inline-python": ("tool_used", "tool: Bash\ninput_match: " + r"""'(?:(?:python\\?"?|\$\{?\w+\}?\\?"?)\s+(?:-c\s*\\?["'']|-\s*<<)|\.py\s*<<|>\s*[^\s|;&]*\.py\b)'""" + "\nmin: 0\nmax: 0", 1,
                         "No inline Python (`python -c`, `\"$PY\" -c`, Python fed on stdin as `python - <<EOF`) and no script "
                         "written through a shell redirect or heredoc: the toolkit is the allowlisted scripts. (Matched "
                         "against the JSON-encoded command, so a quote may carry a backslash.)"),
    "crops-through-crop-sh": ("tool_used", 'tool: Write\ninput_match: "\\\\.construction[\\\\\\\\/]+crops"\nmin: 0\nmax: 0', 1,
                              ".construction/crops/ is written only by crop.sh, never by the Write tool."),
}

def case(name, skill_cmd, prompt, needs="", graders=None, checks=None, extra_setup="", api=False, offline_write=False,
         tools="[Read, Write, Edit, Glob, Grep, Bash, Skill]", own_extraction_ok=False, max_turns=80, timeout=1500,
         desc="", drop=()):
    d = f"agentcm-mode-{name}"
    base_tag = name.replace("-offline", "").replace("-reconcile", "")
    write(f"{d}/prompt.md", f"""
        ---
        name: {d}
        description: {json.dumps(desc)}
        tags: [agentcm-mode, {base_tag}]
        runs: 1
        max_turns: {max_turns}
        timeout_seconds: {timeout}
        allowed_tools: {tools}
        append_system_prompt: {json.dumps(UNATTENDED)}
        ---

        {skill_cmd}

        {prompt}
    """)
    write(f"{d}/case.yaml", 'schema_version: "1.1"\ncontext:\n  scaffold_script: setup.sh\n')
    flag = "--api " if api else ""
    mode = "AgentCM running (the stand-in API)" if api else "AgentCM closed"
    write(f"{d}/setup.sh", f"""
        #!/usr/bin/env bash
        # AgentCM's contract fixture in the workspace, {mode}{' (needs: ' + needs + ')' if needs else ''}.
        set -euo pipefail
        HERE="$(cd "$(dirname "${{BASH_SOURCE[0]}}")" && pwd)"
        bash "$HERE/../agentcm-mode/setup-fixture.sh" {flag}{needs}
        {extra_setup}
        bash "$HERE/../agentcm-mode/baseline.sh"
    """)
    for sub in ("graders", "checks"):  # written in full below: a grader or check dropped here must not linger
        shutil.rmtree(P / d / sub, ignore_errors=True)
    all_graders = dict(COMMON)
    if not own_extraction_ok:
        checks = {"no-own-extraction": (f"t.check_no_own_extraction({READ_MANUALS!r}, {DRAWINGS!r})", 1), **(checks or {})}
    all_graders.update(graders or {})
    for g in drop:
        all_graders.pop(g, None)
    for gname, (gtype, fields, weight, body) in all_graders.items():
        write(f"{d}/graders/{gname}.md", f"---\ntype: {gtype}\n{fields}\nweight: {weight}\n---\n{body}\n")
    harness = "checks:\n  - name: agentcm-mode-writes\n    script: ../_checks/agentcm_mode_writes.py\n    weight: 1\n"
    for cname, (body, weight) in (checks or {}).items():
        # the head at column 0, like a multi-line body: indenting it under the body's first line only left the
        # rest of a multi-line body unindented, and nothing to dedent (IndentationError)
        write(f"{d}/checks/{cname.replace('-', '_')}.py",
              "#!/usr/bin/env python3\n"
              f'"""{cname}: graded from the trace and the requests the stand-in API received."""\n'
              "import sys\nfrom pathlib import Path\n"
              'sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_checks"))\n'
              "import agentcm_trace as t\n\n"
              + textwrap.dedent(body).strip("\n") + "\n")
        harness += f"  - name: {cname}\n    script: checks/{cname.replace('-', '_')}.py\n    weight: {weight}\n"
    if api or offline_write:
        harness += f'allow_urls: ["{STUB}"]\n'
    harness += "teardown: ../agentcm-mode/teardown.sh\n"
    write(f"{d}/harness.yaml", harness)


# Read-only skills, AgentCM closed --------------------------------------------------------------------
sched_list = "; ".join(f"{r[0]} ({r[3]})" for r in SCHEDULES)
case("project-setup", "/construction:project-setup", "",
     desc="AgentCM closed: reports AgentCM's orientation from INDEX.md and changes nothing.",
     graders={
         "no-claude-md": ("file_exists", 'path: "CLAUDE.md"\nexists: false', 2, "The skill writes no CLAUDE.md block in an AgentCM project."),
         "orientation": ("llm", "", 3,
                         f"The final message reports AgentCM's orientation from INDEX.md: {len(SHEETS)} drawing sheets ({', '.join(SHEETS)}), "
                         f"{len(SCHEDULES)} schedule(s) or legend(s) ({sched_list}), {len(SECTIONS)} specification section(s)"
                         + (f", {len(UNREAD)} manual(s) AgentCM could not read" if UNREAD else "")
                         + (f", {ROOMS.group(1)} room(s) known" if ROOMS else "")
                         + ", and contract version 1. It does not recommend splitting drawings or specs and does not inventory the project's files."),
     })
case("sheet-splitter", "/construction:sheet-splitter", "",
     desc="AgentCM closed: splits nothing and points to Split & Rename.",
     graders={
         "no-split": ("file_exists", 'path: "sheets/**/*.pdf"\nexists: false', 2, "No sheet PDF is written."),
         "split-and-rename": ("regex", 'target: trace\npattern: "Split (&|and) Rename"', 2, "The run points the PE to Split & Rename in AgentCM."),
         "reads-sheet-index": ("regex", 'target: trace\npattern: "index[\\\\\\\\/]+sheet_index\\\\.yaml"', 1, "It read AgentCM's sheet index to confirm the split."),
     })
xref_text = "; ".join(f"{x[0]} ({x[4]} \"{x[-1]}\")" for x in XREF_A102)
case("construction-guide", "/construction:construction-guide",
     "From this project's documents: what size is door D103 on the A-601 door schedule, and what refers to sheet A-102? "
     "Look at the drawing to confirm the door's size. Cite where each answer comes from.",
     desc="AgentCM closed: answers from AgentCM's files and checks the drawing through crop.sh with a box from the sheet files.",
     graders={"answer": ("llm", "", 3,
                         f"The answer gives door D103's size as {D103_SIZE} from the A-601 door schedule, and names what refers to A-102 "
                         f"({xref_text}), each with its source in AgentCM's files. A search that finds nothing is reported as not found in "
                         "AgentCM's index, never as proof that something does not exist.")},
     checks={"crop-box": ('t.check_crop_covers("A-601", "D103")', 2)})
case("pe-review", "/construction:pe-review",
     "Review sheets A-101 and A-102 of this project for coordination problems and anything missing. Look at the drawings where the text points.",
     desc="AgentCM closed: reviews from AgentCM's files and looks at the drawings through crop.sh.",
     graders={"review": ("llm", "", 2, "The review works from AgentCM's sheet files for A-101 and A-102, names at least one concrete item on those sheets (a room, the detail 5 reference, a note), cites the sheet for each, and does not claim anything is absent from the project because a search of AgentCM's index found nothing.")},
     checks={"crop-box": ('t.check_crop_from_sheet_files(["A-101", "A-102"])', 2)})
case("rfi-drafter", "/construction:rfi-drafter",
     "Draft an RFI asking the architect to confirm the hardware for door D101: the A-601 door schedule gives its size but the hardware column needs confirming. Research it in the project documents first and look at the schedule on the drawing.",
     desc="AgentCM closed: researches through AgentCM's files and crop.sh, and writes the RFI outside .construction/.",
     graders={"docx": ("file_exists", 'path: "**/RFI-*.docx"', 2, "The RFI is written as a Word document in the project folder.")},
     checks={"crop-box": ('t.check_crop_covers("A-601", "D101")', 2)})
case("bid-tabulator", "/construction:bid-tabulator",
     "Tabulate the five BP-09 Flooring bids in 03 - Subcontractor Bids/BP-09 Flooring/ into a bid comparison workbook named Bid_Tabulation_BP-09.xlsx in the project root.",
     desc="AgentCM closed: takes the project name from project.yaml and keeps its working data under .construction/skills/.",
     extra_setup='bash "$HERE/../_fixtures/prepare-real-workspace.sh" bids',
     own_extraction_ok=True, drop=("agentcm-files",),
     graders={"workbook": ("file_exists", 'path: "Bid_Tabulation_BP-09*.xlsx"', 2, "The bid tab is written to the project root.")})
for name, cmd, prompt, needs, desc, extra in [
    ("submittal-log-generator", "/construction:submittal-log-generator", "Generate the submittal register for this project.", "specs",
     "AgentCM closed: reads AgentCM's specs/ text, and runs spec-splitter only for a manual AgentCM could not read.", ""),
    ("code-researcher", "/construction:code-researcher", "What do the specifications require of the hollow metal doors and frames, and which codes and standards do they cite?", "specs",
     "AgentCM closed: reads AgentCM's specs/ and sheet files and looks through crop.sh.", ""),
    ("bid-evaluator", "/construction:bid-evaluator", "Evaluate the tabulated BP-09 Flooring bids against the specifications.", "specs flooring-specs",
     "AgentCM closed: reads the bid's spec sections from AgentCM's specs/. Waits for spec sections matching the BP-09 bids.",
     'bash "$HERE/../_fixtures/prepare-real-workspace.sh" bids bids-tabulated'),
    ("subcontract-writer", "/construction:subcontract-writer", "Write the BP-09 Flooring subcontract for the awarded bidder from the firm's template.", "specs flooring-specs",
     "AgentCM closed: reads the scope's spec sections from AgentCM's specs/. Waits for spec sections matching the BP-09 bids.",
     'bash "$HERE/../_fixtures/prepare-real-workspace.sh" subcontract-template awarded-bid'),
    ("submittal-review", "/construction:submittal-review", "Review the submittal package in 06 - Submittals/.", "specs pdfs submittal",
     "AgentCM closed: reads the governing sections from AgentCM's specs/ and the trace from tags.tsv and xrefs.tsv. Waits for a submittal against one of the fixture's sections.", ""),
]:
    case(name, cmd, prompt, needs=needs, desc=desc, extra_setup=extra,
         graders={"reads-specs": ("regex", 'target: trace\npattern: "\\\\.construction[\\\\\\\\/]+specs[\\\\\\\\/]+"', 2, "The skill read the spec text from AgentCM's specs/, not its own extraction.")})
unread_names = [r[0] for r in UNREAD]
case("spec-splitter", "/construction:spec-splitter", "",
     needs="specs unread-manual pdfs", own_extraction_ok=True,
     desc="AgentCM closed: splits and extracts only the manual INDEX.md lists as not read; never re-splits one AgentCM read.",
     graders={"says-so": ("regex", 'target: trace\npattern: "(?i)could not read|not read"', 1, "The run says which manual AgentCM could not read and that it splits that one.")},
     checks={"only-unread-manual": (f'''
import re
read = {sorted({r[4] for r in SECTIONS})!r}
unread = {unread_names!r}
work = [c for c in t.bash_commands() if re.search(r"split_spec_manual|extract_spec_text|rasterize_page|pdf_text", c)]
touched_read = [m for m in read if any(Path(m).name in c for c in work)]
touched_unread = [m for m in unread if any(Path(m).name in c for c in work)]
t.finish(not touched_read and bool(touched_unread), f"split and extracted only {{touched_unread}}",
         f"read manuals touched {{touched_read}}; unread manuals touched {{touched_unread}}")
''', 3)})

# Writing skills, AgentCM running (the stand-in) -------------------------------------------------------
case("tag-audit-and-takeoff", "/construction:tag-audit-and-takeoff",
     "Do a quantity takeoff of the room tags on sheets A-101 and A-102 and save its QTO JSON and workbook. "
     "Then run the same takeoff a second time from scratch, as a new session would (read AgentCM's files again and look "
     "at the sheets again), save its own QTO JSON and workbook, and report both runs' confirmed, to confirm and possible "
     "totals. Ingest new detections without asking; nobody can confirm.",
     needs="groups pending pdfs", api=True, max_turns=150, timeout=2700,
     desc="AgentCM running: a takeoff twice gives the same three numbers, confirmed = accepted tags only; a tag sent again is already held, a rejected one is listed once.",
     graders={"three-numbers": ("llm", "", 2, "The final message reports, for each takeoff, confirmed, to confirm and the possible total, with the caveat that confirmed holds only tags accepted in AgentCM's review and the rest wait to be confirmed there; the two takeoffs report the same numbers."),
              "outcomes": ("llm", "", 1, "The final message reports what AgentCM answered for the tags it sent (how many were created, already held, rejected or had no free text), taken from AgentCM's answer, not assumed from the request succeeding. If AgentCM answered that the PE rejected a tag, that tag is listed apart and not counted; a run with no rejected tag passes this on the rest.")},
     checks={"twice-the-same": (f"t.check_takeoff_twice(accepted={ROOM_ACCEPTED}, pending={ROOM_PENDING})", 3),
             "contract-requests": ("t.check_contract_requests()", 1),
             "tag-ingest": ('t.check_tag_ingest(["A-101", "A-102"])', 2),
             "crop-box": ('t.check_crop_from_sheet_files(["A-101", "A-102"])', 1)})
case("schedule-extractor", "/construction:schedule-extractor",
     "Extract the door schedule on A-601 into Excel and ingest it into AgentCM.",
     needs="pdfs", api=True, own_extraction_ok=True,
     desc="AgentCM running: ingests with A-601's id from sheet_index.yaml and the contract's payload.",
     checks={"schedule-ingest": ('t.check_schedule_ingest("A-601")', 3), "contract-requests": ("t.check_contract_requests()", 1)})

case("schedule-extractor-reconcile", "/construction:schedule-extractor",
     "The PE edited the schedule they exported from AgentCM; the workbook is in the project root, named \"... Schedule (PE edits).xlsx\". "
     "Apply their changes back to AgentCM.",
     needs="schedule-id", api=True, own_extraction_ok=True,
     drop=("agentcm-files",),  # the round trip's sources are the workbook's _agentcm_meta and the rows endpoint
     extra_setup='"$HERE/../../../bin/construction-python" "$HERE/../agentcm-mode/make_edited_schedule.py" .',
     desc="AgentCM running: the Excel round trip reads the stored rows from the rows endpoint and reconciles the PE's edit.",
     checks={"reconcile": ('''
sys.path.insert(0, str(t.HERE.parent / "agentcm-mode"))
import make_edited_schedule as m
e = m.expectation()
t.check_reconcile(e["schedule_id"], e["row_key"], e["column"], e["value"])
''', 3), "contract-requests": ("t.check_contract_requests()", 1)})

# Writing skills, AgentCM closed: the payload waits in the outbox ----------------------------------------
OUTBOX_MSG = ("llm", "", 2, "The final message says AgentCM was not running, that what would have been sent waits in the skill's outbox under .construction/skills/, and asks the PE to start AgentCM and rerun the write step. The rest of the task (the deliverable) is done from AgentCM's files.")
case("tag-audit-and-takeoff-offline", "/construction:tag-audit-and-takeoff",
     "Do a quantity takeoff of the room tags on sheet A-101 and save its QTO JSON and workbook. Ingest new detections without asking; nobody can confirm.",
     needs="groups pending pdfs api-url", offline_write=True,
     desc="AgentCM closed: the takeoff runs from the files and the tag ingest waits in the outbox.",
     graders={"tells-the-pe": OUTBOX_MSG,
              "workbook": ("file_exists", 'path: "QTO_*.xlsx"', 1, "The QTO workbook is still produced.")},
     checks={"outbox": ('t.check_outbox("tag-audit-and-takeoff", "tag", "A-101")', 3)})
case("schedule-extractor-offline", "/construction:schedule-extractor",
     "Extract the door schedule on A-601 into Excel and ingest it into AgentCM.",
     needs="pdfs api-url", offline_write=True, own_extraction_ok=True,
     desc="AgentCM closed: the Excel is written and the schedule ingest waits in the outbox.",
     graders={"tells-the-pe": OUTBOX_MSG,
              "workbook": ("file_exists", 'path: "*.xlsx"', 1, "The schedule workbook is still produced.")},
     checks={"outbox": ('t.check_outbox("schedule-extractor", "schedule", "A-601")', 3)})

print("facts:", {"sheets": SHEETS, "schedules": len(SCHEDULES), "sections": len(SECTIONS), "unread": unread_names,
                 "rooms": ROOMS.group(1) if ROOMS else None, "D103": D103_SIZE, "xrefs to A-102": len(XREF_A102),
                 "room tags accepted/pending": (ROOM_ACCEPTED, ROOM_PENDING), "recorded answers": len(RESP)})
print("written:", sorted(p.name for p in P.glob("agentcm-mode*")))
