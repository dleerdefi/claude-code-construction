#!/usr/bin/env python3
"""For the reconcile case: the workbook a PE edits after exporting AgentCM's ingested schedule.

The ingested schedule is the one whose schedules/<sheet>-<n>.tsv comment line carries schedule_id=<uuid>;
its rows are AgentCM's recorded answer to GET .../schedules/<id>/rows. The workbook is written by the
plugin's own schedule_to_xlsx.py (the anchored export: _agentcm_meta, _row_key), then one cell is edited:
the first row's last column other than the row key gets " (REV)" appended. expectation() gives that edit
without writing anything, for the check.

Usage: make_edited_schedule.py <workspace>   writes "<sheet> Schedule (PE edits).xlsx" there
"""
import json
import re
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parents[2] / "skills" / "schedule-extractor" / "scripts"


def ingested():
    """(sheet, schedule_id, rows) of the fixture's ingested schedule."""
    for f in sorted((HERE / "construction" / "schedules").glob("*.tsv")):
        first = f.read_text(encoding="utf-8").splitlines()[0]
        m = re.search(r"schedule_id=([0-9a-f-]{36})", first)
        if m:
            responses = json.loads((HERE / "api-responses.json").read_text(encoding="utf-8"))
            rows = next(v for k, v in responses.items() if k.startswith("GET ") and k.endswith(f"/schedules/{m.group(1)}/rows"))
            return f.stem.rsplit("-", 1)[0], m.group(1), rows
    raise SystemExit("no ingested schedule (schedule_id=<uuid>) in the fixture")


def key_column(rows):
    cols = list(rows[0]["data"].keys())
    return next((c for c in cols if all(str(r["data"].get(c)) == str(r["entityIdentifier"]) for r in rows)), None)


def expectation():
    sheet, sid, rows = ingested()
    key = key_column(rows)
    col = [c for c in rows[0]["data"] if c != key][-1]
    return {"sheet": sheet, "schedule_id": sid, "row_key": str(rows[0]["entityIdentifier"]), "column": col,
            "value": f"{rows[0]['data'].get(col) or ''} (REV)".strip()}


def main(ws):
    from openpyxl import load_workbook
    sys.path.insert(0, str(SCRIPTS))
    from schedule_to_xlsx import schedule_to_xlsx
    sheet, sid, rows = ingested()
    exp = expectation()
    sheet_ids = (HERE / "construction" / "index" / "sheet_index.yaml").read_text(encoding="utf-8")
    sheet_id = re.search(r'id: "([^"]+)"\s+sheet_number: "' + re.escape(sheet) + '"', sheet_ids).group(1)
    out = Path(ws) / f"{sheet} Schedule (PE edits).xlsx"
    with tempfile.TemporaryDirectory() as d:
        data = Path(d) / "rows.json"
        data.write_text(json.dumps([{**{k: v for k, v in r["data"].items()}} for r in rows]), encoding="utf-8")
        schedule_to_xlsx(str(data), output=str(out), schedule_type="schedule", sheet=sheet, schedule_id=sid,
                         sheet_id=sheet_id, row_key_column=key_column(rows) or "")
    wb = load_workbook(out)
    ws_data = next(wb[n] for n in wb.sheetnames if n != "_agentcm_meta")
    header_row = next(r for r in range(1, ws_data.max_row + 1) if ws_data.cell(r, 1).value == "_row_key")
    headers = {ws_data.cell(header_row, c).value: c for c in range(2, ws_data.max_column + 1)}
    row = next(r for r in range(header_row + 1, ws_data.max_row + 1) if str(ws_data.cell(r, 1).value) == exp["row_key"])
    ws_data.cell(row, headers[exp["column"]], exp["value"])
    wb.save(out)
    print(f"edited {out.name}: {exp['row_key']} / {exp['column']} = {exp['value']!r}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
