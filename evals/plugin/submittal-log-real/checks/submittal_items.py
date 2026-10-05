"""Check the submittal log against the Part 1 submittal articles of the three sections.

Each ground-truth item is matched by the leading phrase of its text (e.g. "Product Data",
"Samples for Verification", "Qualification Data") within the log rows for the same section.
"""
import glob
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

WORKSPACE = Path(os.environ["EVAL_WORKSPACE"])
GT = Path(os.environ["EVAL_CASE_DIR"]).parent / "_fixtures" / "sanibel" / "ground_truth" / "submittals.json"


def key(s) -> str:
    return re.sub(r"\D", "", str(s or ""))


def leading_phrase(item: str) -> str:
    return re.split(r"[:.]", item, maxsplit=1)[0].strip().lower()


def main() -> int:
    files = sorted(glob.glob(str(WORKSPACE / "**" / "Submittal_Log*.xlsx"), recursive=True))
    if not files:
        print("no Submittal_Log*.xlsx"); print("score: 0"); return 1
    wb = openpyxl.load_workbook(files[-1], data_only=True)
    ws = wb["Submittal Log"] if "Submittal Log" in wb.sheetnames else wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    header = [str(c or "").strip().lower() for c in rows[0]]
    sec_i = next((i for i, h in enumerate(header) if "spec section" in h or h == "section"), 0)
    desc_i = [i for i, h in enumerate(header) if "description" in h or "type" in h]
    log = {}
    for r in rows[1:]:
        text = " ".join(str(r[i] or "") for i in desc_i).lower()
        log.setdefault(key(r[sec_i]), []).append(text)

    truth = json.loads(GT.read_text(encoding="utf-8"))
    total, matched, missing = 0, 0, []
    for section, data in truth.items():
        for article, items in data["articles"].items():
            for item in items:
                total += 1
                phrase = leading_phrase(item)
                if any(phrase in t for t in log.get(key(section), [])):
                    matched += 1
                else:
                    missing.append(f"{section}: {phrase}")
    recall = matched / total if total else 0
    rows_per_section = {s: len(log.get(key(s), [])) for s in truth}
    print(f"log rows per section: {rows_per_section}; other sections in log: {sorted(k for k in log if k not in {key(s) for s in truth})}")
    print(f"ground-truth items matched: {matched}/{total} (recall {recall:.2f})")
    for m in missing[:10]:
        print("  missing:", m)
    print(f"score: {round(recall, 3)}")
    return 0 if recall >= 0.8 else 1


if __name__ == "__main__":
    sys.exit(main())
