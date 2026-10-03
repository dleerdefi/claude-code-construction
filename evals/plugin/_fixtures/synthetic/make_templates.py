#!/usr/bin/env python3
"""Generate the synthetic "firm template" fixtures for the evals.

Firm: "Example Builders, Inc." -- an invented placeholder general contractor.
No real company name, logo or letterhead is used.

Writes into ./templates/ (next to this script):
  RFI Template - Example Builders.docx          firm RFI form (template mode)
  rfi_template_map.json                         mapping for rfi_export.py
  Subcontract Template - Example Builders.docx  firm subcontract with planted defects
  README.md                                     what is here and how evals use it

Output is deterministic: core properties are pinned and the zip entries are
rewritten with a fixed timestamp, so the sha256 in the mapping is stable.

Run:  bin/construction-python evals/plugin/_fixtures/synthetic/make_templates.py
"""

import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent / "templates"
RFI_NAME = "RFI Template - Example Builders.docx"
SC_NAME = "Subcontract Template - Example Builders.docx"
FOOTER = "EVAL FIXTURE — fictional firm template for a synthetic evaluation"
RFI_MARKER = "EXAMPLE BUILDERS, INC. — REQUEST FOR INFORMATION — FORM EB-RFI-02"
SC_MARKER = "EXAMPLE BUILDERS, INC. STANDARD SUBCONTRACT AGREEMENT — FORM EB-SC-15"
FIXED_TS = datetime(2026, 1, 1, 0, 0, 0)


def new_doc():
    """Blank python-docx document with pinned metadata and the fixture footer."""
    doc = Document()
    cp = doc.core_properties
    cp.author = "Example Builders, Inc. (fictional)"
    cp.last_modified_by = cp.author
    cp.created = cp.modified = FIXED_TS
    cp.title = "Eval fixture"
    cp.comments = FOOTER
    doc.sections[0].footer.paragraphs[0].text = FOOTER
    return doc


def save_deterministic(doc, path):
    """Save, then rewrite the zip with a fixed timestamp for byte-stable output."""
    tmp = path.with_suffix(".tmp")
    doc.save(str(tmp))
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, zin.read(info.filename))
    tmp.unlink()


# ---------------------------------------------------------------------------
# RFI template
# ---------------------------------------------------------------------------

RFI_ROWS = ["Project", "Project No.", "RFI No.", "Date", "From", "To",
            "Subject", "Spec Section", "Drawing Reference"]

RFI_PLACEHOLDERS = [
    ("DESCRIPTION", "[DESCRIPTION]", "description"),
    ("SUGGESTED RESOLUTION", "[SUGGESTED RESOLUTION]", "suggested_resolution"),
    ("IMPACT IF NOT RESOLVED", "[IMPACT IF NOT RESOLVED]", "impact"),
    ("ATTACHMENTS", "[ATTACHMENTS]", "attachments"),
]

# (field, row) for the label/value table; col is always 1 (value column).
RFI_TABLE_FIELDS = [
    ("project_name", 0), ("project_number", 1), ("rfi_number", 2), ("date", 3),
    ("from.company", 4), ("to.company", 5), ("subject", 6),
    ("spec_section", 7), ("drawing_ref", 8),
]


def build_rfi():
    doc = new_doc()
    doc.add_paragraph().add_run(RFI_MARKER).bold = True
    doc.add_paragraph("Submit one RFI per issue. Responses requested within 7 calendar days.")
    table = doc.add_table(rows=len(RFI_ROWS), cols=2)
    table.style = "Table Grid"
    for i, label in enumerate(RFI_ROWS):
        table.cell(i, 0).paragraphs[0].add_run(label).bold = True
        # value cells intentionally left empty
    for heading, placeholder, _ in RFI_PLACEHOLDERS:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(placeholder)
    doc.add_paragraph()
    doc.add_paragraph("Submitted by: ______________________________    Date: ______________")
    path = OUT / RFI_NAME
    save_deterministic(doc, path)
    return path


def write_rfi_map(docx_path):
    mappings = [{"field": f, "type": "table_cell", "table_index": 0, "row": r, "col": 1}
                for f, r in RFI_TABLE_FIELDS]
    mappings += [{"field": f, "type": "placeholder", "search_text": ph}
                 for _, ph, f in RFI_PLACEHOLDERS]
    data = {
        "template_path": docx_path.name,
        "template_hash": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
        "field_mappings": mappings,
    }
    (OUT / "rfi_template_map.json").write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Subcontract template
# ---------------------------------------------------------------------------

BOILER = {
    1: ["The Subcontract Documents consist of this Agreement, the Exhibits listed below, the Prime Contract "
        "between Contractor and Owner (without compensation terms), and the Drawings and Specifications "
        "identified in Exhibit A.",
        "In the event of conflict, the order of precedence is: (1) executed change orders, (2) this "
        "Agreement, (3) the Exhibits, (4) the Prime Contract, (5) the Specifications, (6) the Drawings. "
        "Subcontractor is bound to Contractor as Contractor is bound to Owner."],
    5: ["Contractor shall pay Subcontractor monthly for Work in place, based on an approved Schedule of "
        "Values and an application submitted by the 25th of each month, less retainage of ten percent (10%).",
        "Payment to Subcontractor is conditioned upon receipt of payment from Owner. Subcontractor "
        "shall furnish lien waivers and certified payroll with each application.",
        "Final payment is due after completion of the Work, delivery of closeout documents, and "
        "final waiver of liens."],
    6: ["No change to the Work is valid unless authorized by a written change order or construction "
        "change directive signed by Contractor.",
        "Subcontractor shall submit pricing for a proposed change within seven (7) days of request, "
        "itemized by labor, material, equipment, and markup."],
    7: ["Subcontractor shall maintain commercial general liability, automobile, workers' compensation, and "
        "umbrella coverage in the amounts stated in Exhibit D, and name Contractor and Owner as "
        "additional insureds.",
        "When required by Contractor, Subcontractor shall furnish performance and payment bonds from a "
        "surety acceptable to Contractor, in the amount of the Subcontract Sum.",
        "Certificates of insurance shall be delivered before mobilization."],
    9: ["Subcontractor shall employ competent workers and comply with all applicable labor and "
        "employment laws, wage determinations, and project labor requirements.",
        "Subcontractor shall promptly remove any worker whom Contractor reasonably finds unfit or "
        "disruptive to the Project."],
    10: ["Subcontractor shall indemnify, defend, and hold harmless Contractor, Owner, Architect, and "
         "their officers, directors, and employees from claims, damages, losses, and expenses arising out "
         "of the performance of the Work, except to the extent caused by the negligence of the Contractor.",
         "This obligation survives completion of the Work and termination of this Agreement."],
    12: ["Subcontractor shall comply with the Contractor's site safety program, applicable OSHA "
         "standards, and the Safety Plan in Exhibit E, and shall designate a competent person on site.",
         "Subcontractor shall report all incidents to Contractor's superintendent within 24 hours."],
    13: ["If Subcontractor fails to perform the Work, supply sufficient workers or materials, or comply "
         "with this Agreement, Contractor may give seven (7) days' written notice to cure.",
         "If the default is not cured, Contractor may take over the Work or terminate this Agreement and "
         "charge Subcontractor the resulting excess cost.",
         "Contractor may terminate for convenience on written notice, paying for Work performed to the "
         "date of termination."],
    14: ["The parties shall first attempt to resolve any dispute by direct negotiation between "
         "senior representatives, then by non-binding mediation.",
         "Unresolved disputes shall be decided in the courts of the state where the Project is located. "
         "Subcontractor shall continue the Work during any dispute."],
    15: ["This Agreement is the entire agreement of the parties and may be amended only in writing. "
         "Subcontractor shall not assign this Agreement without Contractor's written consent.",
         "Notices shall be in writing and delivered to the addresses of the parties stated above. "
         "If any provision is held invalid, the remainder stays in effect."],
}

FILL = {
    2: ["Subcontractor shall furnish all labor, materials, equipment, and supervision to perform the "
        "following scope of Work, as further described in Exhibit A:", "<<SCOPE>>"],
    3: ["Subcontractor shall perform the Work in accordance with the following schedule and "
        "Exhibit C, and shall coordinate with Contractor's project schedule:", "[SCHEDULE]"],
    4: ["Contractor shall pay Subcontractor, for performance of the Work, the Subcontract Sum of:",
        "[CONTRACT SUM]", "The Subcontract Sum is based on the Schedule of Values in Exhibit B."],
    8: ["Subcontractor shall submit the following submittals, including product data, shop drawings, "
        "and samples, in time to avoid delay to the Work:", "[SUBMITTALS]"],
    11: ["Subcontractor warrants the Work as follows, in addition to any longer manufacturer warranty:",
         "[WARRANTY]"],
}

ARTICLES = ["Subcontract Documents", "Scope of Work", "Schedule", "Subcontract Sum",
            "Progress Payments", "Changes", "Insurance and Bonds", "Submittals", "Labor",
            "Indemnification", "Warranty", "Safety", "Default and Termination",
            "Dispute Resolution", "General Provisions"]

EXHIBITS = ["Exhibit A — Scope", "Exhibit B — Schedule of Values", "Exhibit C — Schedule",
            "Exhibit D — Insurance Requirements", "Exhibit E — Safety Plan (informational)"]


def build_subcontract():
    doc = new_doc()
    doc.add_paragraph().add_run(SC_MARKER).bold = True
    doc.add_paragraph("This Subcontract Agreement is made as of [DATE] between Example Builders, Inc. "
                      "(\"Contractor\") and [SUBCONTRACTOR] (\"Subcontractor\").")
    parties = doc.add_table(rows=4, cols=2)
    parties.style = "Table Grid"
    for i, (k, v) in enumerate([("Date", "[DATE]"), ("Subcontractor", "[SUBCONTRACTOR]"),
                                ("Subcontract No.", "[SUBCONTRACT NUMBER]"),
                                ("Subcontract Sum", "[CONTRACT SUM]")]):
        parties.cell(i, 0).text = k
        parties.cell(i, 1).text = v
    doc.add_paragraph("Example from prior project: Harbor View Office Building, "
                      "Subcontractor: Oldco Drywall, Inc., Subcontract No. EB-21001")
    for n, title in enumerate(ARTICLES, 1):
        doc.add_heading(f"Article {n} — {title}", level=1)
        for text in BOILER.get(n) or FILL[n]:
            doc.add_paragraph(text)
    doc.add_paragraph("IN WITNESS WHEREOF, the parties have executed this Agreement as of the date "
                      "first written above.")
    sig = doc.add_table(rows=3, cols=2)
    sig.style = "Table Grid"
    for i, (a, b) in enumerate([("CONTRACTOR: Example Builders, Inc.", "SUBCONTRACTOR: [SUBCONTRACTOR]"),
                                ("By: ____________________", "By: ____________________"),
                                ("Name/Title: ______________", "Name/Title: ______________")]):
        sig.cell(i, 0).text = a
        sig.cell(i, 1).text = b
    doc.add_heading("Exhibits", level=1)
    for ex in EXHIBITS:
        doc.add_paragraph(ex, style="List Bullet")
    path = OUT / SC_NAME
    save_deterministic(doc, path)
    return path


README = f"""# Synthetic firm templates (eval fixtures)

Generated by `make_templates.py` (deterministic; rerun to regenerate). The firm,
"Example Builders, Inc.", is an invented placeholder. Both documents carry the footer
"{FOOTER}".
All names, numbers and text are invented; no real company, logo or letterhead is used.

## Files

| File | What it is |
|---|---|
| `RFI Template - Example Builders.docx` | Firm RFI form. Marker paragraph `{RFI_MARKER}`. 2-column label/value table (9 rows, empty value cells) plus single-paragraph placeholders `[DESCRIPTION]`, `[SUGGESTED RESOLUTION]`, `[IMPACT IF NOT RESOLVED]`, `[ATTACHMENTS]`. It deliberately has no "RESPONSE (to be completed by design professional)" block, so graders can tell template mode from generic mode. |
| `rfi_template_map.json` | Mapping for `skills/rfi-drafter/scripts/rfi_export.py` (table_cell for the 9 table rows, placeholder for the 4 body fields) with the sha256 of the docx. |
| `Subcontract Template - Example Builders.docx` | Firm subcontract. Marker `{SC_MARKER}`, parties block with `[DATE]`, `[SUBCONTRACTOR]`, `[SUBCONTRACT NUMBER]`, `[CONTRACT SUM]`, 15 articles, signature block, Exhibits A-E. |

## RFI mapping choices

- `From` and `To` are single cells, so `from.company` and `to.company` are mapped to them.
  The person names (`from.name`, `to.name`) are NOT mapped and are omitted from the output.
- Placeholders are in body paragraphs (not headers/footers), one paragraph each.

## Planted subcontract defects

1. Article 10 carve-out reads "except to the extent caused by the negligence of the Contractor".
   The word "sole" is deliberately missing; the skill must fix it ("sole negligence") and flag it for review.
2. Article 5 is pay-when-paid ("payment to Subcontractor is conditioned upon receipt of payment
   from Owner") with no time-certain outside date; the skill should flag it.
3. The parties block carries stale example data: "Example from prior project: Harbor View Office
   Building, Subcontractor: Oldco Drywall, Inc., Subcontract No. EB-21001". None of it
   may appear in a new subcontract.

Article classes: boilerplate (1, 5, 6, 7, 9, 10, 12, 13, 14, 15) is preserved; articles 2, 3, 4, 8, 11
hold placeholders (`<<SCOPE>>`, `[SCHEDULE]`, `[CONTRACT SUM]`, `[SUBMITTALS]`, `[WARRANTY]`) to be
generated or filled.

## How the evals use them

- rfi-drafter: the template and mapping are supplied; graders check the output contains the
  EB-RFI-02 marker, filled table cells, no leftover `[` placeholders, and no generic RESPONSE block.
- subcontract-writer: graders check the output keeps the EB-SC-15 structure, replaces all placeholders,
  contains "sole negligence" in Article 10 (with a legal flag), flags pay-when-paid, and contains none of
  "Harbor View", "Oldco Drywall", "EB-21001".
"""


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rfi = build_rfi()
    write_rfi_map(rfi)
    sc = build_subcontract()
    (OUT / "README.md").write_text(README, encoding="utf-8")
    for p in (rfi, OUT / "rfi_template_map.json", sc, OUT / "README.md"):
        print(f"{p.stat().st_size:>8}  {p}")


if __name__ == "__main__":
    main()
